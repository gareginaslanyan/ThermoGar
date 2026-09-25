"""21-L, step 2: defaults and 10B candidates, one fresh app process per run.

Every run: free memory >= 3.0 GiB, a new app process on port 8639 with an
empty state root %TEMP%\\tg21l_state\\<run number> (a new user: no
prepare_state.py), light theme, window 1440x900, the sidebar expanded (as in
the 21-B frames). The database is chosen in the sidebar and the sidebar
composition is compared with DATABASE_DEFINITIONS of the app. «Умолчание» —
the primary button of the screen is pressed with nothing changed;
«кандидат» — only the listed fields are changed, then the button is pressed.
Limit: 15 min from the press; past it the app is closed (taskkill /T /F on
the process tree of the port) and the run is recorded as not finished.

Frames — the whole page (sidebar included) after the result or the refusal,
parts ch1, ch2… as in the 21-B scripts (kadry.capture):
results/wave21_l/10b/<screen>_<base>_<variant>[_<step>][_chN].png.
Per run: results/wave21_l/10b/runs/<nn>.json (what was changed, time,
alerts, captions, metrics, tables, charts); app log: 10b/logi/<nn>_app.log.

    D:\\Pets\\ThermoGar\\.venv-windows\\Scripts\\python.exe -B -X utf8 results\\wave21_l\\scripts\\progon_10b.py [run ids…]
"""

from __future__ import annotations

import json
import os
import re
import subprocess
import sys
import time
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import kadry  # noqa: E402  (21-I copy: role="radio" for the switches)

mgs = kadry.mgs
REPO = kadry.REPO_ROOT
OUT = REPO / "results" / "wave21_l" / "10b"
RUNS_DIR = OUT / "runs"
LOGS = OUT / "logi"
PORT = 8639
PY = r"D:\Pets\ThermoGar\.venv-windows\Scripts\python.exe"
STATE_BASE = Path(os.environ["TEMP"]) / "tg21l_state"
LIMIT_S = 15 * 60
NO_FRAMES = False
BASES = {"ni": mgs.DB_NI, "al": mgs.DB_AL, "fe": mgs.DB_FE}
DEFAULTS = {  # app/ThermoGar_app.py DATABASE_DEFINITIONS
    "ni": ("Ni", "атомные %", "Al=15"),
    "fe": ("Fe", "массовые %", "C=0.20, Cr=11.5, Ni=0.7"),
    "al": ("Al", "атомные %", "Cu=4"),
}

SCREENS = {
    "S01": (["Расчёты", "Одна температура"], "Рассчитать равновесие"),
    "S02": (["Расчёты", "Температурный диапазон"], "Построить график по температуре"),
    "S03": (["Расчёты", "Изменение состава"], "Построить график по составу"),
    "S05": (["Диаграммы", "Многокомпонентное T–X"], "Построить многокомпонентное сечение"),
    "S06": (["Диаграммы", "Тройная при T = const"], "Построить тройную диаграмму"),
    "S07": (["Диаграммы", "Карта доли фазы"], "Построить карту доли фазы"),
    "S08": (["Затвердевание"], "Рассчитать затвердевание"),
    "S09": (["Энергии", "Энергии фаз"], "Рассчитать энергии фаз"),
    "S10": (["Энергии", "Движущая сила"], "Рассчитать движущую силу"),
    "S11": (["Энергии", "T₀"], "Рассчитать T₀"),
    "S12": (["Свойства", "Плотность"], "Рассчитать плотность и объёмные доли"),
    "S13": (["Свойства", "Плотность по T"], "Построить плотность по температуре"),
    "S14": (["Свойства", "Упругие свойства"], "Получить фазовые доли"),
    "S15": (["Свойства", "Вклады упрочнения"], "Рассчитать вклады"),
    "S16": (["Кинетика", "Диффузия и гомогенизация", "Однофазная пара"],
            "Рассчитать однофазную диффузию"),
    "S17": (["Кинетика", "Диффузия и гомогенизация", "Многофазная гомогенизация"],
            "Рассчитать гомогенизацию"),
    "S19": (["Кинетика", "Выделения"], "Рассчитать кинетику выделений"),
}
VRH_BUTTON = "Рассчитать Voigt–Reuss–Hill"


# --- field actions: (kind, label, value) ---------------------------------------------------

def num(label, value):
    return ("number", label, value)


def select(label, value):
    return ("select", label, value)


def text(label, value):
    return ("text_area", label, value)


def check(label, value=True):
    return ("checkbox", label, value)


def multi(label, values):
    return ("multiselect", label, values)


def run(rid, screen, base, variant, changes=(), *, step="", note="", before_frame=False,
        vrh=False):
    return {"id": rid, "screen": screen, "base": base, "variant": variant,
            "changes": list(changes), "step": step, "note": note,
            "before_frame": before_frame, "vrh": vrh}


K2_S02 = [num("От, °C", "500"), num("До, °C", "900"), num("Шаг, °C", "100")]
K2_S13 = [num("Температура от, °C", "500"), num("Температура до, °C", "900"),
          num("Шаг температуры, °C", "100")]
SOURCE_21L = "Проверка умолчаний 21-Л"
CONFIRM = "Подтверждаю область применимости введённых коэффициентов"

RUNS = [
    run("01", "S02", "fe", "umolch"),
    run("02", "S02", "fe", "K2", K2_S02),
    run("03", "S13", "fe", "umolch"),
    run("04", "S13", "fe", "K2", K2_S13),
    run("05", "S03", "fe", "umolch"),
    run("06", "S03", "fe", "K1", [select("Изменяемый элемент", "C"),
                                  num("C: от, мас.%", "0.1"), num("C: до, мас.%", "0.5"),
                                  num("C: шаг, мас.%", "0.1")]),
    run("07", "S03", "al", "umolch"),
    run("08", "S03", "al", "K1", [select("Изменяемый элемент", "Cu"),
                                  num("Cu: от, ат.%", "1"), num("Cu: до, ат.%", "5"),
                                  num("Cu: шаг, ат.%", "1")]),
    run("09", "S03", "ni", "umolch"),
    run("10", "S02", "ni", "umolch"),
    run("11", "S02", "al", "umolch"),
    run("12", "S01", "al", "umolch"),
    run("13", "S05", "ni", "umolch"),
    run("14", "S05", "ni", "K3", [text("Постоянные добавки, ат.%", "Cr=8"),
                                  num("Al: от, ат.%", "0"), num("Al: до, ат.%", "30"),
                                  num("Шаг по составу, ат.%", "2"),
                                  num("Температура от, °C", "626.85"),
                                  num("Температура до, °C", "1626.85"),
                                  num("Шаг по температуре, °C", "50")]),
    run("15", "S05", "al", "umolch"),
    run("16", "S05", "fe", "umolch"),
    run("17", "S06", "al", "umolch"),
    run("18", "S06", "fe", "umolch"),
    run("19", "S07", "al", "umolch"),
    run("20", "S07", "fe", "umolch"),
    run("21", "S08", "ni", "umolch"),
    run("22", "S08", "al", "umolch"),
    run("23", "S08", "fe", "umolch"),
    run("24", "S09", "al", "umolch"),
    run("25", "S09", "fe", "umolch"),
    run("26", "S10", "ni", "umolch"),
    run("27", "S10", "al", "umolch"),
    run("28", "S10", "fe", "umolch"),
    run("29", "S11", "al", "umolch"),
    run("30", "S11", "fe", "umolch"),
    # 31, 32: К-4, only if 29 / 30 refuse or find no T₀.
    run("31", "S11", "al", "K4", [num("Нижняя граница поиска T₀, °C", "530"),
                                  num("Верхняя граница поиска T₀, °C", "830")]),
    run("32", "S11", "fe", "K4", [num("Нижняя граница поиска T₀, °C", "600"),
                                  num("Верхняя граница поиска T₀, °C", "760")]),
    run("33", "S11", "ni", "dop", [num("Нижняя граница поиска T₀, °C", "1230"),
                                   num("Верхняя граница поиска T₀, °C", "1630"),
                                   select("Фаза 1", "FCC_A1"),
                                   select("Фаза 2", "GAMMA_PRIME")], step="gp"),
    run("34", "S11", "ni", "dop", [num("Нижняя граница поиска T₀, °C", "1230"),
                                   num("Верхняя граница поиска T₀, °C", "1630"),
                                   select("Фаза 1", "FCC_A1"),
                                   select("Фаза 2", "LIQUID")], step="liq"),
    run("35", "S12", "al", "umolch"),
    run("36", "S13", "al", "umolch"),
    run("37", "S14", "al", "umolch", step="shag1"),
    run("38", "S14", "ni", "umolch", step="shag2", vrh=True),
    run("39", "S15", "ni", "umolch", step="b", before_frame=True),
    run("40", "S15", "ni", "dop", [text("Источник и область применимости входов", SOURCE_21L)],
        step="v"),
    run("41", "S15", "ni", "dop", [text("Источник и область применимости входов", SOURCE_21L),
                                   check(CONFIRM)], step="g"),
    run("42", "S15", "al", "dop", [text("Источник и область применимости входов", SOURCE_21L),
                                   check(CONFIRM)], step="g"),
    run("43", "S15", "fe", "dop", [text("Источник и область применимости входов", SOURCE_21L),
                                   check(CONFIRM)], step="g"),
    run("44", "S16", "al", "umolch"),
    run("45", "S16", "fe", "umolch"),
    run("46", "S17", "ni", "umolch"),
    run("47", "S17", "ni", "K5", [multi("Фазы локального равновесия", ["FCC_A1", "NIAL"])]),
    run("48", "S17", "fe", "umolch"),
    run("49", "S19", "ni", "umolch"),
    # 50: К-6, only if 49 does not finish in 15 min.
    run("50", "S19", "ni", "K6", [num("Время выдержки, ч", "1")]),
    run("51", "S19", "al", "umolch"),
    run("52", "S19", "fe", "umolch"),
]


# --- app process ------------------------------------------------------------------------------

def port_pids(port: int) -> set[int]:
    out = subprocess.run(["netstat", "-ano", "-p", "TCP"], capture_output=True, text=True,
                         encoding="oem", errors="replace").stdout
    pids = set()
    for line in out.splitlines():
        parts = line.split()
        if len(parts) >= 5 and parts[1].endswith(f":{port}") and parts[3] == "LISTENING":
            pids.add(int(parts[4]))
    return pids


def kill_tree(pids) -> list[str]:
    said = []
    for pid in sorted(pids):
        result = subprocess.run(["taskkill", "/PID", str(pid), "/T", "/F"],
                                capture_output=True, text=True, encoding="oem",
                                errors="replace")
        said.append((result.stdout + result.stderr).strip())
    return said


def start_app(rid: str):
    # A repeated attempt of a run (script fix) takes a new empty folder <id>_<n>.
    state = STATE_BASE / rid
    attempt = 1
    while state.exists():
        attempt += 1
        state = STATE_BASE / f"{rid}_{attempt}"
    state.mkdir(parents=True)
    env = dict(os.environ, THERMOGAR_STATE_ROOT=str(state), PYTHONHASHSEED="0",
               PYTHONUTF8="1")
    log = open(LOGS / f"{state.name}_app.log", "wb")
    process = subprocess.Popen(
        [PY, "-m", "streamlit", "run", r"app\ThermoGar_app.py", "--server.headless", "true",
         "--server.port", str(PORT), "--server.address", "127.0.0.1",
         "--client.toolbarMode", "auto"],
        cwd=str(REPO), env=env, stdout=log, stderr=subprocess.STDOUT,
        creationflags=subprocess.CREATE_NEW_PROCESS_GROUP,
    )
    deadline = time.monotonic() + 180
    while not mgs.app_is_ready(PORT):
        if process.poll() is not None:
            raise RuntimeError(f"app exited with {process.returncode}")
        if time.monotonic() > deadline:
            raise RuntimeError("app did not start in 180 s")
        time.sleep(1)
    return process, log, state


def stop_app(process, log) -> list[str]:
    pids = port_pids(PORT) | ({process.pid} if process.poll() is None else set())
    said = kill_tree(pids)
    try:
        process.wait(timeout=30)
    except subprocess.TimeoutExpired:
        said.append("process did not exit in 30 s")
    log.close()
    deadline = time.monotonic() + 30
    while port_pids(PORT) and time.monotonic() < deadline:
        time.sleep(1)
    if port_pids(PORT):
        said.append(f"port {PORT} still busy: {port_pids(PORT)}")
    return said


# --- page helpers -----------------------------------------------------------------------------

SIDEBAR_VALUES_JS = """
() => {
  const side = document.querySelector('[data-testid="stSidebar"]');
  const byLabel = (testid, label) => [...side.querySelectorAll(`[data-testid="${testid}"]`)]
    .find(e => (e.querySelector('[data-testid="stWidgetLabel"]') || {}).innerText?.trim() === label);
  const base = byLabel('stSelectbox', 'База материалов');
  const bal = byLabel('stSelectbox', 'Элемент-основа');
  const units = byLabel('stRadio', 'Единицы состава');
  const add = byLabel('stTextArea', 'Добавки');
  const selected = (el, label) => {
    const input = el.querySelector('input');
    const shownText = el.innerText.replace(label, '').trim();
    return shownText || (input ? (input.value || input.getAttribute('aria-label') || '') : '');
  };
  const r = side.getBoundingClientRect();
  return {
    database: base ? selected(base, 'База материалов') : null,
    balance: bal ? selected(bal, 'Элемент-основа') : null,
    units: units ? [...units.querySelectorAll('input[type="radio"]')].filter(i => i.checked)
      .map(i => i.closest('label').innerText.trim())[0] : null,
    additions: add ? add.querySelector('textarea').value : null,
    sidebar_width: Math.round(r.width), sidebar_right: Math.round(r.right),
  };
}
"""

# Everything the main area shows below (and alerts above) the pressed button.
RESULT_JS = """
(buttonText) => {
  const main = document.querySelector('[data-testid="stMain"]');
  const shown = e => e.offsetParent !== null && e.getBoundingClientRect().height > 0;
  const cut = (s, n) => (s || '').replace(/\\s+/g, ' ').trim().slice(0, n);
  const buttons = [...main.querySelectorAll('button')].filter(shown)
    .filter(b => b.innerText.trim() === buttonText);
  const anchor = buttons.length ? buttons[buttons.length - 1] : null;
  const after = el => !anchor || (anchor.compareDocumentPosition(el) & Node.DOCUMENT_POSITION_FOLLOWING);
  const out = {button_found: !!anchor, button_disabled: anchor ? anchor.disabled : null,
               alerts: [], texts: [], metrics: [], tables: [], charts: 0, switches: [],
               expanders: []};
  for (const el of main.querySelectorAll('[data-testid="stAlert"]')) {
    if (!shown(el)) continue;
    const c = el.querySelector('[data-testid^="stAlertContent"]');
    out.alerts.push({kind: c ? c.dataset.testid.replace('stAlertContent', '').toLowerCase() : '?',
                     where: after(el) ? 'below' : 'above', text: cut(el.innerText, 2000)});
  }
  for (const el of main.querySelectorAll('[data-testid="stException"]')) {
    if (shown(el)) out.alerts.push({kind: 'exception', where: after(el) ? 'below' : 'above',
                                    text: cut(el.innerText, 2000)});
  }
  for (const el of main.querySelectorAll(
      '[data-testid="stMarkdownContainer"], [data-testid="stCaptionContainer"], [data-testid="stHeading"]')) {
    if (!shown(el) || !after(el) || el.closest('[data-testid="stAlert"]') ||
        el.closest('[data-testid="stWidgetLabel"]') || el.closest('[data-testid="stExpander"] summary') ||
        el.closest('button') || el.closest('[data-testid="stMetric"]')) continue;
    const t = cut(el.innerText, 1500);
    if (t) out.texts.push(t);
  }
  for (const el of main.querySelectorAll('[data-testid="stMetric"]')) {
    if (shown(el) && after(el)) out.metrics.push(cut(el.innerText, 200));
  }
  for (const el of main.querySelectorAll('[data-testid="stDataFrame"], [data-testid="stTable"]')) {
    if (!shown(el) || !after(el)) continue;
    const grid = el.querySelector('[role="grid"], table');
    const heads = [...el.querySelectorAll('[role="columnheader"], th')].map(h => cut(h.innerText, 60))
      .filter(Boolean);
    out.tables.push({rows: grid ? Number(grid.getAttribute('aria-rowcount') || 0) - 1 : null,
                     cols: grid ? Number(grid.getAttribute('aria-colcount') || 0) : null,
                     headers: heads.slice(0, 40),
                     // Rendered rows of the grid (glide renders the visible part only).
                     cells: [...el.querySelectorAll('[role="row"]')].slice(0, 80)
                       .map(r => [...r.querySelectorAll('[role="gridcell"], td')]
                         .map(c => cut(c.innerText, 40)))
                       .filter(r => r.length)});
  }
  out.charts = [...main.querySelectorAll(
    '[data-testid="stImage"] img, [data-testid="stPlotlyChart"], [data-testid="stVegaLiteChart"]')]
    .filter(e => shown(e) && after(e)).length;
  out.switches = [...main.querySelectorAll('[role="radiogroup"][aria-label]')]
    .filter(g => shown(g) && after(g)).map(g => g.getAttribute('aria-label'));
  out.expanders = [...main.querySelectorAll('[data-testid="stExpander"] summary')]
    .filter(e => shown(e) && after(e)).map(e => cut(e.innerText.replace(/keyboard_arrow_\\w+/g, ''), 80));
  return out;
}
"""


# All rows of the result grids: the glide grid renders the visible rows only, so the
# scroller is moved down step by step and rows are collected by aria-rowindex.
GRID_ROWS_JS = """
async (buttonText) => {
  const main = document.querySelector('[data-testid="stMain"]');
  const shown = e => e.offsetParent !== null && e.getBoundingClientRect().height > 0;
  const sleep = ms => new Promise(r => setTimeout(r, ms));
  const buttons = [...main.querySelectorAll('button')].filter(shown)
    .filter(b => b.innerText.trim() === buttonText);
  const anchor = buttons.length ? buttons[buttons.length - 1] : null;
  const after = el => !anchor || (anchor.compareDocumentPosition(el) & Node.DOCUMENT_POSITION_FOLLOWING);
  const out = [];
  for (const el of main.querySelectorAll('[data-testid="stDataFrame"]')) {
    if (!shown(el) || !after(el)) continue;
    const grid = el.querySelector('[role="grid"]');
    const scroller = el.querySelector('.dvn-scroller');
    const expander = el.closest('[data-testid="stExpander"]');
    const title = expander ? expander.querySelector('summary').innerText
      .replace(/keyboard_arrow_\w+/g, '').trim() : null;
    const heads = [...el.querySelectorAll('[role="columnheader"]')].map(h => h.innerText.trim());
    const rows = new Map();
    const take = () => {
      for (const r of el.querySelectorAll('[role="row"]')) {
        const i = Number(r.getAttribute('aria-rowindex'));
        const cells = [...r.querySelectorAll('[role="gridcell"]')].map(c => c.innerText.trim());
        if (cells.length) rows.set(i, cells);
      }
    };
    take();
    if (scroller) {
      for (let top = 0; top <= scroller.scrollHeight + 400 && rows.size < 5000; top += 200) {
        scroller.scrollTop = top;
        await sleep(60);
        take();
      }
      scroller.scrollTop = 0;
    }
    out.push({expander: title, rowcount: grid ? Number(grid.getAttribute('aria-rowcount')) - 1 : null,
              headers: heads, rows: [...rows.entries()].sort((a, b) => a[0] - b[0]).map(e => e[1])});
  }
  return out;
}
"""


GRID_TAKE_JS = """
(args) => {
  const [buttonText, index] = args;
  const main = document.querySelector('[data-testid="stMain"]');
  const shown = e => e.offsetParent !== null && e.getBoundingClientRect().height > 0;
  const buttons = [...main.querySelectorAll('button')].filter(shown)
    .filter(b => b.innerText.trim() === buttonText);
  const anchor = buttons.length ? buttons[buttons.length - 1] : null;
  const after = el => !anchor || (anchor.compareDocumentPosition(el) & Node.DOCUMENT_POSITION_FOLLOWING);
  const grids = [...main.querySelectorAll('[data-testid="stDataFrame"]')].filter(e => shown(e) && after(e));
  if (index === null) return grids.length;
  const el = grids[index];
  const grid = el.querySelector('[role="grid"]');
  const expander = el.closest('[data-testid="stExpander"]');
  const rows = [];
  for (const r of el.querySelectorAll('[role="row"]')) {
    const cells = [...r.querySelectorAll('[role="gridcell"]')].map(c => c.innerText.trim());
    if (cells.length) rows.push([Number(r.getAttribute('aria-rowindex')), cells]);
  }
  return {expander: expander ? expander.querySelector('summary').innerText
            .replace(/keyboard_arrow_\w+/g, '').trim() : null,
          rowcount: grid ? Number(grid.getAttribute('aria-rowcount')) - 1 : null,
          headers: [...el.querySelectorAll('[role="columnheader"]')].map(h => h.innerText.trim()),
          rows};
}
"""


def read_grids(page, button: str) -> list[dict]:
    """All rows of the result grids: the mouse wheel over a grid scrolls it (glide renders
    the visible rows only); rows are collected by aria-rowindex."""

    out = []
    count = page.evaluate(GRID_TAKE_JS, [button, None])
    grids = mgs.main_area(page).locator('[data-testid="stDataFrame"]:visible')
    for index in range(count):
        data = page.evaluate(GRID_TAKE_JS, [button, index])
        rows = {i: cells for i, cells in data["rows"]}
        # The index-th grid after the button is the (total - count + index)-th visible grid.
        locator = grids.nth(grids.count() - count + index)
        locator.scroll_into_view_if_needed()
        box = locator.bounding_box()
        if box and data["rowcount"] and len(rows) < data["rowcount"]:
            page.mouse.move(box["x"] + box["width"] / 2, box["y"] + min(box["height"] / 2, 200))
            stale = 0
            while len(rows) < data["rowcount"] and stale < 6:
                page.mouse.wheel(0, 300)
                page.wait_for_timeout(120)
                before = len(rows)
                rows.update({i: c for i, c in page.evaluate(GRID_TAKE_JS, [button, index])["rows"]})
                stale = stale + 1 if len(rows) == before else 0
            page.mouse.move(5, 5)
        data["rows"] = [rows[i] for i in sorted(rows)]
        out.append(data)
    return out


def download_excel(page, stem: str) -> dict | None:
    """Repeat runs only: the screen's own «Скачать Excel» -> 10b/tablicy/<stem>.xlsx;
    boundary lines and nodes are counted from the sheet «Границы», and the maximum of
    every «… доля …» column is taken (full tables, not the rendered part of the grid)."""

    button = mgs.main_area(page).get_by_role("button", name="Скачать Excel", exact=True)
    if not button.count():
        return None
    folder = OUT / "tablicy"
    folder.mkdir(exist_ok=True)
    with page.expect_download(timeout=120_000) as info:
        button.last.click()
    target = folder / f"{stem}.xlsx"
    info.value.save_as(str(target))
    mgs.wait_idle(page)
    import openpyxl

    book = openpyxl.load_workbook(target, read_only=True, data_only=True)
    out: dict = {"file": target.name, "sheets": book.sheetnames, "max": {}}
    for sheet in book.worksheets:
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            continue
        head = [str(h) if h is not None else "" for h in rows[0]]
        if sheet.title == "Границы" and "Тип" in head and "Группа" in head:
            kind, group = head.index("Тип"), head.index("Группа")
            phase = next((i for i, h in enumerate(head) if h.startswith("Фаза")), None)
            lines, nodes, phases = set(), set(), set()
            for row in rows[1:]:
                if str(row[kind]).startswith("Граница"):
                    lines.add(row[group])
                    if phase is not None:
                        phases.add(row[phase])
                elif str(row[kind]).startswith("Инвариант"):
                    nodes.add(row[group])
            out["boundaries"] = {"rows": len(rows) - 1, "lines": len(lines), "nodes": len(nodes),
                                 "line_phases": sorted(map(str, phases))}
        for i, h in enumerate(head):
            if "доля" in h:
                values = [r[i] for r in rows[1:] if isinstance(r[i], (int, float))]
                if values:
                    out["max"][f"{sheet.title}: {h}"] = max(values)
    book.close()
    return out


def open_result_expanders(page) -> list[str]:
    """Boundary tables of the diagrams sit in closed expanders: open them (after the frame)."""

    opened = []
    root = mgs.main_area(page)
    for title in ("Таблица рассчитанных границ и узлов", "Таблица рассчитанных границ"):
        blocks = root.locator('[data-testid="stExpander"]:visible').filter(has_text=title)
        if blocks.count():
            mgs.open_expander(page, root, title)
            opened.append(title)
            break
    return opened


def boundary_counts(grids: list[dict]) -> dict | None:
    for grid in grids:
        heads = grid["headers"]
        if "Тип" not in heads or "Группа" not in heads:
            continue
        kind, group = heads.index("Тип"), heads.index("Группа")
        phase = next((i for i, h in enumerate(heads) if h.startswith("Фаза")), None)
        lines, nodes, phases = set(), set(), set()
        for row in grid["rows"]:
            if len(row) <= max(kind, group):
                continue
            if row[kind].startswith("Граница"):
                lines.add(row[group])
                if phase is not None and len(row) > phase:
                    phases.add(row[phase])
            elif row[kind].startswith("Инвариант"):
                nodes.add(row[group])
        return {"rows_read": len(grid["rows"]), "rowcount": grid["rowcount"],
                "lines": len(lines), "nodes": len(nodes), "line_phases": sorted(phases)}
    return None


def expand_sidebar(page) -> None:
    expand = page.locator('[data-testid="stExpandSidebarButton"]')
    if expand.count() and expand.first.is_visible():
        expand.first.click()
        mgs.wait_idle(page)
    page.wait_for_timeout(500)


def pick_option(page, root, testid: str, label: str, option: str) -> None:
    control = mgs.widget(root, testid, label)
    field = control.locator('input[role="combobox"]').first if testid == "stSelectbox" \
        else control.locator("input").first
    field.scroll_into_view_if_needed()
    field.click()
    if testid == "stMultiSelect":
        field.fill(option)
        page.wait_for_timeout(300)
    page.get_by_role("option").first.wait_for()
    options = page.get_by_role("option", name=option, exact=True)
    if not options.count():
        names = page.get_by_role("option").all_inner_texts()
        raise RuntimeError(f"option {option!r} not in {label!r}: {names}")
    options.first.click()
    mgs.wait_idle(page)


def apply(page, change) -> None:
    kind, label, value = change
    root = mgs.main_area(page)
    if kind == "number":
        mgs.set_number(root, label, value)
    elif kind == "select":
        current = mgs.widget(root, "stSelectbox", label).inner_text().replace(label, "").strip()
        if current != value:
            pick_option(page, root, "stSelectbox", label, value)
    elif kind == "text_area":
        mgs.set_text_area(root, label, value)
    elif kind == "checkbox":
        mgs.set_checkbox(root, label, value)
    elif kind == "multiselect":
        control = mgs.widget(root, "stMultiSelect", label)
        clear = control.locator('[aria-label="Clear all"], [title="Clear all"]')
        if clear.count():
            clear.first.click()
            mgs.wait_idle(page)
        for item in value:
            pick_option(page, root, "stMultiSelect", label, item)
        page.keyboard.press("Escape")
    mgs.wait_idle(page)


def read_field(page, change) -> str:
    kind, label, _ = change
    root = mgs.main_area(page)
    testid = {"number": "stNumberInput", "select": "stSelectbox", "text_area": "stTextArea",
              "checkbox": "stCheckbox", "multiselect": "stMultiSelect"}[kind]
    control = mgs.widget(root, testid, label)
    # A field whose label follows another field (e.g. «C: от» before «C» is chosen) is not there yet.
    if not root.locator(f'[data-testid="{testid}"]:visible').filter(has_text=label).count():
        return None
    if kind == "number":
        return control.locator("input").input_value()
    if kind == "text_area":
        return control.locator("textarea").input_value()
    if kind == "checkbox":
        return str(control.locator('input[type="checkbox"]').is_checked())
    return control.inner_text().replace(label, "").replace("\n", " ").strip()


def frame(page, stem: str) -> list[str]:
    page.evaluate(
        "() => { window.scrollTo(0, 0); for (const e of document.querySelectorAll('*'))"
        " if (e.scrollTop) e.scrollTop = 0; }"
    )
    paths, info = kadry.capture(page, stem, with_sidebar=True)
    return [p.name for p in paths]


def press(page, name: str) -> tuple[float, bool]:
    """Press the button; seconds until the script is idle, and whether it was."""

    root = mgs.main_area(page)
    button = root.get_by_role("button", name=name, exact=True).last
    button.scroll_into_view_if_needed()
    if button.is_disabled():
        return 0.0, True
    started = time.monotonic()
    button.click()
    mgs.CALC_TIMEOUT_MS = LIMIT_S * 1000
    try:
        mgs.wait_idle(page)
    except TimeoutError:
        return time.monotonic() - started, False
    return time.monotonic() - started, True


def do_run(browser, spec: dict) -> dict:
    rid, screen, base = spec["id"], spec["screen"], spec["base"]
    stem = f"{screen}_{base}_{spec['variant']}" + (f"_{spec['step']}" if spec["step"] else "")
    record = {**{k: spec[k] for k in ("id", "screen", "base", "variant", "step", "note")},
              "changes": [list(c) for c in spec["changes"]], "stem": stem}
    free = kadry.free_gib()
    record["free_gib_before"] = round(free, 2)
    print(f"[{rid}] {stem}: free {free:.2f} GiB", flush=True)
    if free < kadry.MIN_FREE_GIB:
        record["status"] = "STOP: free memory below 3.0 GiB"
        return record
    if port_pids(PORT):
        record["status"] = f"STOP: port {PORT} busy before start"
        return record
    t0 = time.monotonic()
    process, log, state = start_app(rid)
    record["state_root"] = str(state)
    record["app_start_s"] = round(time.monotonic() - t0, 1)
    context = browser.new_context(viewport=dict(kadry.VIEWPORT), device_scale_factor=1,
                                  color_scheme="light", locale="ru-RU",
                                  timezone_id="Europe/Moscow", accept_downloads=True)
    context.set_default_timeout(mgs.UI_TIMEOUT_MS)
    page = context.new_page()
    kadry.KADRY = OUT
    try:
        page.goto(f"http://127.0.0.1:{PORT}/", wait_until="domcontentloaded")
        page.get_by_role("tab", name="Расчёты", exact=True).wait_for(timeout=600_000)
        mgs.CALC_TIMEOUT_MS = 600_000
        mgs.wait_idle(page)
        if kadry.current_theme(page) != "Light":
            raise RuntimeError(f"not light: {page.evaluate(kadry.APP_BG_JS)}")
        expand_sidebar(page)
        mgs.set_database(page, BASES[base])
        mgs.wait_idle(page)
        side = page.evaluate(SIDEBAR_VALUES_JS)
        want = DEFAULTS[base]
        side["default_ok"] = (
            side["balance"] == want[0] and side["units"] == want[1]
            and re.sub(r"\s", "", (side["additions"] or "")).lower()
            == re.sub(r"\s", "", want[2]).lower()
        )
        record["sidebar"] = side
        print(f"  sidebar {side}", flush=True)
        if not side["default_ok"]:
            record["status"] = "STOP: sidebar composition is not the default of the database"
            frame(page, f"{stem}_sboy_sostav")
            return record
        path, button = SCREENS[screen]
        for name in path:
            mgs.open_tab(page, name)
        record["fields_before"] = {c[1]: read_field(page, c) for c in spec["changes"]}
        for change in spec["changes"]:
            apply(page, change)
        record["fields_after"] = {c[1]: read_field(page, c) for c in spec["changes"]}
        print(f"  fields {record['fields_before']} -> {record['fields_after']}", flush=True)
        frames: list[str] = []
        if spec["before_frame"]:
            frames += frame(page, f"{screen}_{base}_{spec['variant']}_a")
            record["before"] = page.evaluate(RESULT_JS, button)
        if spec["vrh"]:
            seconds, done = press(page, button)
            record["step1_s"] = round(seconds, 1)
            record["step1"] = page.evaluate(RESULT_JS, button)
            if not done:
                raise TimeoutError("step 1 not finished")
            button = VRH_BUTTON
        record["button"] = button
        seconds, done = press(page, button)
        record["calc_s"] = round(seconds, 1)
        record["finished"] = done
        print(f"  pressed «{button}»: {seconds:.1f} s, finished {done}", flush=True)
        if done:
            page.wait_for_timeout(800)
            record["result"] = page.evaluate(RESULT_JS, button)
            if not NO_FRAMES:
                frames += frame(page, stem)
            if NO_FRAMES:
                record["excel"] = download_excel(page, stem)
            record["expanders_opened"] = open_result_expanders(page)
            record["grids"] = read_grids(page, button)
            record["boundaries"] = boundary_counts(record["grids"])
            record["status"] = "ok"
        else:
            record["status"] = "не завершён за 15 мин"
            try:
                record["result"] = page.evaluate(RESULT_JS, button)
                frames += frame(page, f"{stem}_nezavershen")
            except Exception as error:  # noqa: BLE001
                record["frame_error"] = f"{type(error).__name__}: {error}"
        record["frames"] = frames
        record["free_gib_after"] = round(kadry.free_gib(), 2)
    except Exception as error:  # noqa: BLE001
        record["status"] = f"script error: {type(error).__name__}: {str(error).splitlines()[0]}"
        print(f"  {record['status']}", flush=True)
        try:
            page.screenshot(path=str(OUT / f"_sboy_{rid}_{stem}.png"))
        except Exception:  # noqa: BLE001
            pass
    finally:
        context.close()
        record["taskkill"] = stop_app(process, log)
    return record


def main() -> int:
    global NO_FRAMES
    wanted = [a for a in sys.argv[1:] if not a.startswith("--")]
    # --povtor: the same run again only to read the full result tables; no frames,
    # the record goes to runs/<id>_povtor.json.
    NO_FRAMES = any(a.startswith("--povtor") for a in sys.argv)
    # --povtor2 …: a later repeat writes runs/<id>_povtor2.json, the earlier record stays.
    povtor = next((a[2:] for a in sys.argv if a.startswith("--povtor")), "")
    specs = [r for r in RUNS if not wanted or r["id"] in wanted]
    for folder in (OUT, RUNS_DIR, LOGS):
        folder.mkdir(parents=True, exist_ok=True)
    if port_pids(PORT):
        print(f"STOP: port {PORT} busy", file=sys.stderr)
        return 2
    from playwright.sync_api import sync_playwright

    with sync_playwright() as driver:
        browser = driver.chromium.launch(headless=True)
        try:
            for spec in specs:
                record = do_run(browser, spec)
                suffix = f"_{povtor}" if NO_FRAMES else ""
                (RUNS_DIR / f"{spec['id']}{suffix}.json").write_text(
                    json.dumps(record, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
                print(f"  -> {record['status']}", flush=True)
                if record["status"].startswith("STOP"):
                    return 3
        finally:
            browser.close()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
