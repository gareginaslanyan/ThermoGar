"""Решения владельца по стали 26.09.2026 и решение 12Б, внедрённые в 21-О.

* BL-66 и 2В: T₀ стали по умолчанию — окно 200–950 °C, T₀ найдено в 18
  точках из 21, состав — из заданной сетки во всех строках; окно
  300–1700 °C — 15 из 21, состав тоже во всех строках; «Энергии фаз» и
  «Движущая сила» стали — прежние 300–1700 °C;
* 1Б: «Выделения» стали — «Время выдержки, ч» = 0.01;
* 3В: текст остановки по составу — «ушла ниже нуля» при отрицательной доле,
  прежнее число при остальных значениях;
* 12Б: символы элементов в заголовках выгрузок Excel и CSV (Ni, Al, Cu);
  листы raw затвердевания и «Исходные данные» пакета — без изменений;
  исходные таблицы после записи файла не меняются.
"""

from __future__ import annotations

import io
import sys
import zipfile
from pathlib import Path
from typing import Any, Iterator

import pandas as pd
import pytest

ROOT = Path(__file__).resolve().parent.parent
APP = ROOT / "app"
TOOLS = ROOT / "tools"
for path in (APP, TOOLS):
    if str(path) not in sys.path:
        sys.path.insert(0, str(path))

import streamlit as st  # noqa: E402

# Пакетный расчёт гоняется тем же приложением и тем же приватным профилем,
# что и в test_ui_h (фикстуры app и patched_app_source).
from test_ui_h import (  # noqa: E402,F401
    app,
    batch_csv,
    patched_app_source,
    stored_artifacts,
    widget,
)

KNOWN_FOREIGN_ERRORS = ("BINDING_STALE",)
TZERO_COUNTER = "T₀ найдено в {found} точках из {total}. "


@pytest.fixture(autouse=True)
def _private_state_root(tmp_path, monkeypatch) -> None:
    monkeypatch.setenv("THERMOGAR_STATE_ROOT", str(tmp_path / "state"))


@pytest.fixture(autouse=True)
def _bounded_memory(monkeypatch) -> Iterator[None]:
    """Два воркера на тест, как в test_ui_f; после теста пул закрыт."""

    import gc

    import matplotlib.pyplot as plt
    import thermogar_parallel_ui

    monkeypatch.setattr(thermogar_parallel_ui, "_WORKER_COUNT", 2)
    yield
    thermogar_parallel_ui.close_shared_engines()
    plt.close("all")
    gc.collect()


@pytest.fixture(autouse=True)
def _restore_download_button() -> Iterator[None]:
    download_button = st.download_button
    yield
    st.download_button = download_button


def _start(base: str, state: dict[str, Any] | None = None):
    """Приложение на базе по умолчанию и байты всех кнопок скачивания."""

    from streamlit.testing.v1 import AppTest

    downloads: dict[str, bytes] = {}
    original = st.download_button

    def capture(label: Any, data: Any = None, *args: Any, **kwargs: Any) -> Any:
        name = str(kwargs.get("file_name") or label)
        if isinstance(data, (bytes, bytearray)):
            downloads[name] = bytes(data)
        elif isinstance(data, str):
            downloads[name] = data.encode("utf-8")
        return original(label, data, *args, **kwargs)

    st.download_button = capture
    app_test = AppTest.from_file(str(APP / "ThermoGar_app.py"), default_timeout=1800)
    app_test.session_state["thermogar_database_key"] = base
    for key, value in (state or {}).items():
        app_test.session_state[key] = value
    app_test.run()
    _assert_clean(app_test)
    return app_test, downloads


def _assert_clean(app_test) -> None:
    assert not app_test.exception, [str(e.value) for e in app_test.exception]
    errors = [
        e.value for e in app_test.error
        if not any(known in e.value for known in KNOWN_FOREIGN_ERRORS)
    ]
    assert not errors, errors


def _numbers(app_test) -> dict[str, Any]:
    return {item.key: item.value for item in app_test.number_input if item.key}


def _sheet_header(payload: bytes, sheet: str) -> list[Any]:
    import openpyxl

    book = openpyxl.load_workbook(io.BytesIO(payload), read_only=True)
    return [cell.value for cell in next(book[sheet].iter_rows(max_row=1))]


# --------------------------------------------------------------------------- #
# 3В: текст остановки
# --------------------------------------------------------------------------- #


def test_composition_stop_note_below_zero() -> None:
    import thermogar_precipitation as precipitation

    note = precipitation._composition_stop_note(3.885, "TI", -9.864e-5)
    assert note == (
        "Расчёт остановлен на 3.885 с модельного времени (0.001079 ч): доля Ti "
        "в матрице ушла ниже нуля, баланс масс нарушен. Показана часть расчёта "
        f"до остановки. {precipitation.KWN_COMPOSITION_STOP_CAUSE}"
    )


def test_composition_stop_note_zero_keeps_the_number() -> None:
    import thermogar_precipitation as precipitation

    note = precipitation._composition_stop_note(2.011, "NB", 0.0)
    assert "стала 0 ат.%" in note
    assert "ниже нуля" not in note


# --------------------------------------------------------------------------- #
# 1Б, 2В и BL-66: умолчания и T₀ стали
# --------------------------------------------------------------------------- #


def test_steel_windows_and_holding_time() -> None:
    app_test, _ = _start("fe")
    values = _numbers(app_test)
    assert (values["tzero_t_min_fe"], values["tzero_t_max_fe"]) == (200.0, 950.0)
    assert (values["energy_t_min_fe"], values["energy_t_max_fe"]) == (300.0, 1700.0)
    assert (values["driving_t_min_fe"], values["driving_t_max_fe"]) == (300.0, 1700.0)
    durations = [
        item.value for item in app_test.number_input
        if item.label == "Время выдержки, ч" and item.key.startswith("precipitation_fe_")
    ]
    assert durations == [0.01], durations


@pytest.mark.parametrize(
    ("window", "found"),
    [(None, 18), ((300.0, 1700.0), 15)],
    ids=["200-950", "300-1700"],
)
def test_tzero_steel(window: tuple[float, float] | None, found: int) -> None:
    state = {}
    if window is not None:
        state = {"tzero_t_min_fe": window[0], "tzero_t_max_fe": window[1]}
    app_test, downloads = _start("fe", state)
    app_test.button(key="tzero_calculate").click().run()
    _assert_clean(app_test)

    data = app_test.session_state["tzero_result"]["data"]
    assert len(data) == 21
    composition = data["C, мас.%"]
    assert not composition.isna().any()
    assert composition.tolist() == pytest.approx([0.1 * index for index in range(21)])
    assert int(data["Решение найдено"].sum()) == found
    counter = TZERO_COUNTER.format(found=found, total=21)
    assert any(item.value.startswith(counter) for item in app_test.warning), [
        item.value for item in app_test.warning
    ]
    assert _sheet_header(downloads["ThermoGar_T0.xlsx"], "T0")[0] == "C, мас.%"


# --------------------------------------------------------------------------- #
# 12Б: заголовки выгрузок
# --------------------------------------------------------------------------- #


def test_helper_leaves_raw_and_bare_columns() -> None:
    from thermogar_user_errors import element_columns_for_display

    raw = pd.DataFrame(
        {"T": [1.0], "X(FCC_A1,CR)": [0.1], "X(LIQUID,NI)": [0.2], "NP(FCC_A1)": [0.5]}
    )
    source = pd.DataFrame({"Название": ["A"], "CR": [10.0], "NI": [8.0], "C": [0.2]})
    for table in (raw, source):
        assert list(element_columns_for_display(table).columns) == list(table.columns)


def test_single_equilibrium_excel_ni() -> None:
    app_test, downloads = _start("ni")
    app_test.button(key="single_calculate").click().run()
    _assert_clean(app_test)
    display = app_test.session_state["_thermogar_vlb_b3_result_equilibrium_single"]["display"]
    before = list(display["phase_at"].columns)

    header = _sheet_header(downloads["ThermoGar_equilibrium.xlsx"], "Составы фаз ат")
    assert "Ni, ат.%" in header and "Al, ат.%" in header, header
    assert "NI, ат.%" not in header
    assert list(display["phase_at"].columns) == before
    assert "NI, ат.%" in before


def test_concentration_csv_al() -> None:
    app_test, downloads = _start("al")
    app_test.button(key="concentration_calculate").click().run()
    _assert_clean(app_test)
    data = app_test.session_state["_thermogar_vlb_b3_result_equilibrium_composition_scan"][
        "display"
    ]["data"]

    csv = pd.read_csv(io.BytesIO(downloads["ThermoGar_concentration_scan.csv"]))
    assert csv.columns[0] == "Cu, атомные %"
    assert data.columns[0] == "CU, атомные %"
    header = _sheet_header(downloads["ThermoGar_concentration_scan.xlsx"], "Концентрационный расчёт")
    assert header[0] == "Cu, атомные %"


def test_precipitation_excel_ni() -> None:
    app_test, downloads = _start(
        "ni",
        {
            "precipitation_ni_user_duration_h": 0.001,
            "precipitation_ni_user_bins": 40,
        },
    )
    app_test.button(key="precipitation_ni_user_calculate").click().run()
    assert not app_test.exception
    result = app_test.session_state["thermogar_precipitation_result"]
    before = list(result.matrix_composition.columns)
    app_test.segmented_control(key="precipitation_result_view").set_value(
        "Экспорт и ограничения"
    ).run()

    payload = downloads["ThermoGar_precipitation_GAMMA_PRIME.xlsx"]
    matrix = _sheet_header(payload, "Состав матрицы")
    assert "Al, матрица, ат.%" in matrix and "Ni, матрица, ат.%" in matrix, matrix
    interface = _sheet_header(payload, "Межфазные составы")
    assert "Al, матрица на границе, ат.%" in interface, interface
    assert "Al, выделение на границе, ат.%" in interface, interface
    assert list(result.matrix_composition.columns) == before
    assert "AL, матрица, ат.%" in before


@pytest.mark.slow
def test_solidification_zip_and_raw_sheets_ni() -> None:
    app_test, downloads = _start("ni")
    app_test.button(key="solidification_calculate").click().run()
    _assert_clean(app_test)
    state = app_test.session_state["solidification_result"]
    app_test.segmented_control(key="solidification_result_view").set_value("Выгрузка").run()

    payload = downloads["ThermoGar_solidification.xlsx"]
    for method_key in state["results"]:
        short = "Равновес" if method_key == "equilibrium" else "Scheil"
        raw = state["raw_tables"][method_key]
        assert _sheet_header(payload, f"{short} raw") == list(raw.columns)
        liquid = _sheet_header(payload, f"{short} расплав")
        assert "Ni, ат.%" in liquid, liquid
        assert "NI, ат.%" in state["liquid_tables"][method_key].columns

    with zipfile.ZipFile(io.BytesIO(downloads["ThermoGar_solidification_results.zip"])) as archive:
        for method_key in state["results"]:
            member = pd.read_csv(archive.open(f"{method_key}_liquid_composition.csv"))
            assert "Ni, ат.%" in member.columns, list(member.columns)


@pytest.mark.slow
def test_batch_export_ni(app) -> None:
    import openpyxl

    at, state_root = app()
    uploader = [item for item in at.file_uploader if item.label == "Файл составов"]
    uploader[0].set_value(("batch.csv", batch_csv(("ni",)), "text/csv"))
    at.run()
    widget(at.button, "batch_calculate_button").click()
    at.run()
    result = at.session_state["workspace_batch_result"]["display"]
    before = {name: list(frame.columns) for name, frame in result.items()}

    widget(at.button, "batch_result_export_prepare").click()
    at.run()
    book = openpyxl.load_workbook(
        io.BytesIO(stored_artifacts(state_root)["batch-result-xlsx-v1"]), read_only=True
    )

    def header(sheet: str) -> list[Any]:
        return [cell.value for cell in next(book[sheet].iter_rows(max_row=1))]

    atomic = header("Составы фаз ат")
    assert "Ni, ат.%" in atomic and "Al, ат.%" in atomic, atomic
    assert "Ni, мас.%" in header("Составы фаз мас")
    for name in ("Сводка", "Фазовые доли", "Исходные данные"):
        assert header(name) == before[name], name
    assert {name: list(frame.columns) for name, frame in result.items()} == before
    assert "NI, ат.%" in before["Составы фаз ат"]
