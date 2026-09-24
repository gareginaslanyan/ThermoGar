"""20-А (BL-57, шаг 0): эталон вкладок приложения через AppTest.

Снимает то, что приложение показывает и отдаёт, на сценариях матриц UI
(``tools/test_ui_f.py``, ``tools/test_ui_g.py``, ``tools/test_ui_h.py``) и на
составах 15-Ш, чтобы каждый шаг разреза ``app/ThermoGar_app.py`` можно было
сверить с эталоном побайтово (``tools/tab_snapshot_compare.py``).

Каждый случай — отдельный процесс со своей папкой состояния
(``THERMOGAR_STATE_ROOT``); пул воркеров — ``UI_TEST_POOL_WORKERS`` из
``tools/test_ui_f.py``; байты выгрузок — перехватом ``st.download_button``,
как в ``tools/test_ui_f.py``. Проверок (assert) здесь нет: инструмент
записывает, а не судит.

Что пишется в ``<вывод>/<случай>/``:

* ``karkas.txt`` — элементы экрана по порядку (боковая панель, затем основная
  часть) после запуска и после каждого действия, с путём вложенности;
* ``ekran/tNNN.csv`` — таблицы экрана (одна копия на содержимое, номер —
  по первому появлению);
* ``sostoyanie/<прил>.json`` и ``sostoyanie/<прил>/*.csv`` — session_state
  после последнего шага, таблицы — CSV;
* ``vygruzki/`` — выгрузки: CSV, JSON, PNG и прочее — байты как есть; XLSX —
  список листов и листы в CSV; ZIP и NPZ — состав архива и члены по тем же
  правилам; ``vygruzki.txt`` — перечень пойманных выгрузок и шаг, на котором
  каждая появилась впервые;
* ``texts.txt`` — предупреждения, ошибки, сообщения по порядку;
* ``meta.json`` — случай, шаги, версии пакетов, время;
* ``sha256.txt`` — по всем файлам каталога.

Числа с плавающей точкой везде пишутся через ``repr`` (до последнего бита).

Запуск (из корня дерева):

    python -B -X utf8 tools/tab_snapshot.py list
    python -B -X utf8 tools/tab_snapshot.py run --out results/wave20_a/run1 \\
        --state results/wave20_a/state/run1 --time-csv results/wave20_a/run1_time.csv
    python -B -X utf8 tools/tab_snapshot.py case <случай> --out <каталог> --state <каталог>

``run`` запускает случаи по одному: порог входа — свободная память
``--min-free`` (3,0 ГиБ), аварийный — 1,0 ГиБ по ходу случая.
"""

from __future__ import annotations

import argparse
import csv
import dataclasses
import datetime as dt
import enum
import hashlib
import importlib.metadata
import io
import json
import os
import re
import subprocess
import sys
import time
import zipfile
from pathlib import Path
from typing import Any, Callable

PROJECT_ROOT = Path(__file__).resolve().parent.parent
APP_DIR = PROJECT_ROOT / "app"
APP_PATH = APP_DIR / "ThermoGar_app.py"
TOOLS_DIR = PROJECT_ROOT / "tools"
GIB = 2**30
ABORT_FREE_GIB = 1.0
APP_TIMEOUT_S = 3600
PACKAGES = (
    "streamlit", "pandas", "numpy", "scipy", "matplotlib", "openpyxl",
    "pycalphad", "kawin", "scheil", "symengine", "psutil",
)


# ---------------------------------------------------------------------------
# Представление значений
# ---------------------------------------------------------------------------


def sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def safe_name(text: str) -> str:
    """Имя файла из ключа или имени листа: без символов, запрещённых в Windows."""

    cleaned = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", str(text)).strip(" .")
    return cleaned[:120] or "_"


def cell_text(value: Any) -> str:
    """Ячейка CSV: float — repr, остальное — строкой."""

    import numpy as np

    if value is None:
        return ""
    if isinstance(value, (bool, np.bool_)):
        return str(bool(value))
    if isinstance(value, (float, np.floating)):
        return repr(float(value))
    if isinstance(value, (int, np.integer)):
        return str(int(value))
    if isinstance(value, (list, tuple, dict, set, frozenset)):
        return json.dumps(plain(value), ensure_ascii=False)
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        return value.isoformat()
    return str(value)


def plain(value: Any) -> Any:
    """Значение для JSON в одну строку (подписи, значения полей)."""

    import numpy as np

    if value is None or isinstance(value, (str, bool)):
        return value
    if isinstance(value, np.bool_):
        return bool(value)
    if isinstance(value, (int, np.integer)):
        return int(value)
    if isinstance(value, (float, np.floating)):
        return repr(float(value))
    if isinstance(value, (list, tuple)):
        return [plain(item) for item in value]
    if isinstance(value, (set, frozenset)):
        return sorted(repr(item) for item in value)
    if isinstance(value, dict):
        return {str(key): plain(item) for key, item in value.items()}
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        return value.isoformat()
    if isinstance(value, enum.Enum):
        return str(value.value)
    if hasattr(value, "name") and hasattr(value, "size") and hasattr(value, "getvalue"):
        return {"файл": value.name, "байт": value.size}
    return f"<{type(value).__module__}.{type(value).__qualname__}>"


def frame_csv(frame: Any) -> bytes:
    """DataFrame в CSV: заголовок, индекс (если не 0..n-1), float — repr."""

    import pandas as pd

    if isinstance(frame, pd.Series):
        frame = frame.to_frame()
    index = frame.index
    default_index = isinstance(index, pd.RangeIndex) and index.start == 0 and index.step == 1
    out = io.StringIO(newline="")
    writer = csv.writer(out, lineterminator="\n")
    header = [] if default_index else [cell_text(name) for name in index.names]
    header += [cell_text(column) for column in frame.columns]
    writer.writerow(header)
    for position, row in enumerate(frame.itertuples(index=False, name=None)):
        prefix = []
        if not default_index:
            label = index[position]
            prefix = [cell_text(part) for part in (label if isinstance(label, tuple) else (label,))]
        writer.writerow(prefix + [cell_text(item) for item in row])
    return out.getvalue().encode("utf-8")


# ---------------------------------------------------------------------------
# Запись одного случая
# ---------------------------------------------------------------------------


class Recorder:
    """Всё, что снимается с приложений одного случая."""

    def __init__(self, out: Path) -> None:
        self.out = out
        self.karkas: list[str] = []
        self.texts: list[str] = []
        self.steps: list[dict[str, Any]] = []
        self.screen_tables: dict[str, bytes] = {}
        self.table_names: dict[bytes, str] = {}
        self.downloads: dict[str, bytes] = {}
        self.first_seen: dict[str, int] = {}
        self.current_step = 0
        self.exceptions = 0

    # -- экран --------------------------------------------------------------

    def snapshot(self, app_name: str, at: Any, action: str, seconds: float) -> None:
        self.current_step += 1
        step = self.current_step
        self.steps.append({"шаг": step, "прил": app_name, "действие": action,
                           "секунд": round(seconds, 2)})
        self.karkas.append(f"=== шаг {step} | прил {app_name} | {action}")
        self.texts.append(f"=== шаг {step} | прил {app_name} | {action}")
        for area, block in (("боковая", at.sidebar), ("основная", at.main)):
            self._walk(block, [area])

    def _walk(self, block: Any, path: list[str]) -> None:
        from streamlit.testing.v1.element_tree import Block

        for index in sorted(block.children):
            node = block.children[index]
            if isinstance(node, Block):
                segment = block_segment(node, index)
                self._walk(node, path + [segment] if segment else path)
            else:
                self._element(node, path)

    def _element(self, node: Any, path: list[str]) -> None:
        kind = node.type
        where = " / ".join(path)
        fields: list[str] = []
        proto = node.proto

        def has(name: str) -> bool:
            try:
                return proto.DESCRIPTOR.fields_by_name.get(name) is not None
            except AttributeError:
                return False

        if has("label"):
            fields.append(f"подпись={json.dumps(proto.label, ensure_ascii=False)}")
        try:
            key = node.key
        except Exception:  # noqa: BLE001
            key = None
        if key:
            fields.append(f"ключ={json.dumps(key, ensure_ascii=False)}")
        if has("disabled"):
            fields.append(f"disabled={bool(proto.disabled)}")
        if has("help") and proto.help:
            fields.append(f"подсказка={json.dumps(proto.help, ensure_ascii=False)}")
        if has("placeholder") and proto.placeholder:
            fields.append(f"заполнитель={json.dumps(proto.placeholder, ensure_ascii=False)}")
        if has("options") and kind in ("selectbox", "radio", "multiselect", "select_slider", "button_group"):
            fields.append(f"варианты={json.dumps(list(proto.options), ensure_ascii=False)}")
        for bound in ("min", "max", "step"):
            if kind == "number_input" and has(bound):
                flag = {"min": "has_min", "max": "has_max"}.get(bound)
                if flag is None or getattr(proto, flag):
                    fields.append(f"{bound}={repr(float(getattr(proto, bound)))}")
        text: Any = None
        if kind in ("markdown", "caption", "latex", "divider", "error", "warning", "info",
                    "success", "title", "header", "subheader", "text", "code", "toast"):
            text = node.value
        elif kind == "json":
            text = proto.body
        elif kind == "metric":
            text = {"значение": proto.body, "дельта": proto.delta}
        elif kind == "exception":
            self.exceptions += 1
            text = {"тип": proto.type, "сообщение": proto.message}
        if text is not None:
            fields.append(f"текст={json.dumps(plain(text), ensure_ascii=False)}")
        if kind in ("dataframe", "table"):
            try:
                frame = node.value
                data = frame_csv(frame)
                # Порядковое имя по первому появлению содержимого, а не хэш: метка
                # времени внутри таблицы не должна менять имя файла.
                name = self.table_names.get(data)
                if name is None:
                    name = f"ekran/t{len(self.table_names) + 1:03d}.csv"
                    self.table_names[data] = name
                    self.screen_tables[name] = data
                fields.append(f"столбцы={json.dumps([str(c) for c in frame.columns], ensure_ascii=False)}")
                fields.append(f"строк={len(frame)}")
                fields.append(f"таблица={name}")
            except Exception as error:  # noqa: BLE001
                fields.append(f"таблица=<не прочитана: {type(error).__name__}>")
        elif kind == "imgs":
            fields.append(f"картинок={len(proto.imgs)}")
        elif kind not in ("button", "download_button", "form_submit_button", "empty") and kind not in (
            "markdown", "caption", "latex", "divider", "error", "warning", "info", "success",
            "title", "header", "subheader", "text", "code", "toast", "json", "metric", "exception",
        ):
            try:
                value = node.value
            except Exception as error:  # noqa: BLE001
                value = f"<нет значения: {type(error).__name__}>"
            fields.append(f"значение={json.dumps(plain(value), ensure_ascii=False)}")
        line = f"{where} | {kind}" + ("" if not fields else " | " + " | ".join(fields))
        self.karkas.append(line)
        if kind in ("error", "warning", "info", "success", "toast", "exception"):
            self.texts.append(f"{kind} | {where} | {json.dumps(plain(text), ensure_ascii=False)}")

    # -- выгрузки ------------------------------------------------------------

    def capture_download(self, name: str, data: Any) -> None:
        if callable(data):
            payload = b"<callable>"
        elif isinstance(data, (bytes, bytearray)):
            payload = bytes(data)
        elif isinstance(data, str):
            payload = data.encode("utf-8")
        elif hasattr(data, "getvalue"):
            payload = bytes(data.getvalue())
        else:
            payload = repr(type(data)).encode("utf-8")
        if name not in self.first_seen:
            self.first_seen[name] = self.current_step + 1
        self.downloads[name] = payload

    # -- запись на диск -----------------------------------------------------

    def write(self, apps: dict[str, Any], meta: dict[str, Any]) -> None:
        out = self.out
        files: dict[str, bytes] = {}
        files["karkas.txt"] = ("\n".join(self.karkas) + "\n").encode("utf-8")
        files["texts.txt"] = ("\n".join(self.texts) + "\n").encode("utf-8")
        files.update(self.screen_tables)
        for app_name, at in apps.items():
            files.update(session_files(app_name, at))
        # Байты и суммы выгрузок здесь не пишутся: XLSX и ZIP несут дату внутри,
        # а сами выгрузки лежат в vygruzki/ и сравниваются там.
        log = ["файл | впервые на шаге"]
        log += [f"{name} | {step}" for name, step in self.first_seen.items()]
        files["vygruzki.txt"] = ("\n".join(log) + "\n").encode("utf-8")
        for name, payload in self.downloads.items():
            files.update(download_files(f"vygruzki/{safe_name(name)}", payload))
        meta = dict(meta)
        meta["шаги"] = self.steps
        meta["исключений_на_экране"] = self.exceptions
        meta["файлов"] = len(files) + 2
        files["meta.json"] = json.dumps(meta, ensure_ascii=False, indent=1).encode("utf-8")
        for rel, data in files.items():
            target = out / rel
            target.parent.mkdir(parents=True, exist_ok=True)
            target.write_bytes(data)
        lines = [f"{sha256(files[rel])}  {rel}" for rel in sorted(files)]
        (out / "sha256.txt").write_bytes(("\n".join(lines) + "\n").encode("utf-8"))


def block_segment(node: Any, index: int) -> str | None:
    kind = node.type
    if kind == "tab":
        return f"вкладка[{node.label}]"
    if kind in ("expander", "status"):
        return f"раскрытие[{node.label}]"
    if kind == "column":
        return f"колонка#{index}"
    if kind == "form":
        return f"форма[{node.proto.form.form_id}]"
    if kind in ("popover", "dialog", "chat_message"):
        return f"{kind}#{index}"
    return None


def download_files(base: str, payload: bytes) -> dict[str, bytes]:
    """Выгрузка на диск: XLSX — листы в CSV, ZIP/NPZ — члены, прочее — байты."""

    lower = base.lower()
    if lower.endswith((".xlsx", ".xlsm")):
        return xlsx_files(base, payload)
    if lower.endswith((".zip", ".npz")) and payload[:2] == b"PK":
        files: dict[str, bytes] = {}
        with zipfile.ZipFile(io.BytesIO(payload)) as archive:
            members = [info for info in archive.infolist() if not info.is_dir()]
            # Только имена по порядку: размер и CRC члена-XLSX меняются с датой внутри,
            # содержимое членов сравнивается отдельными файлами.
            listing = [m.filename for m in members]
            files[f"{base}.sostav.txt"] = ("\n".join(listing) + "\n").encode("utf-8")
            for member in members:
                name = "/".join(safe_name(part) for part in member.filename.split("/"))
                files.update(download_files(f"{base}__zip/{name}", archive.read(member)))
        return files
    return {base: payload}


def xlsx_files(base: str, payload: bytes) -> dict[str, bytes]:
    import openpyxl

    book = openpyxl.load_workbook(io.BytesIO(payload), data_only=False)
    files = {f"{base}.listy.txt": ("\n".join(book.sheetnames) + "\n").encode("utf-8")}
    for number, sheet in enumerate(book.worksheets, start=1):
        out = io.StringIO(newline="")
        writer = csv.writer(out, lineterminator="\n")
        for row in sheet.iter_rows(values_only=True):
            writer.writerow([cell_text(value) for value in row])
        files[f"{base}__{number:02d}_{safe_name(sheet.title)}.csv"] = out.getvalue().encode("utf-8")
    return files


def session_files(app_name: str, at: Any) -> dict[str, bytes]:
    """session_state: JSON со значениями, таблицы — отдельными CSV."""

    import pandas as pd

    tables: dict[str, bytes] = {}
    prefix = f"sostoyanie/{safe_name(app_name)}"

    def convert(value: Any, path: str, depth: int, seen: set[int]) -> Any:
        import numpy as np

        if isinstance(value, (pd.DataFrame, pd.Series)):
            data = frame_csv(value)
            name = f"{prefix}/{safe_name(path)}.csv"
            suffix = 1
            while name in tables and tables[name] != data:
                suffix += 1
                name = f"{prefix}/{safe_name(path)}__{suffix}.csv"
            tables[name] = data
            columns = [str(c) for c in (value.columns if isinstance(value, pd.DataFrame) else [value.name])]
            return {"__таблица__": name, "строк": len(value), "столбцы": columns}
        if value is None or isinstance(value, (str, bool, int, float)) and not isinstance(value, enum.Enum):
            return value
        if isinstance(value, np.generic):
            return value.item()
        if isinstance(value, (bytes, bytearray, memoryview)):
            data = bytes(value)
            return {"__байты__": len(data), "sha256": sha256(data)}
        if isinstance(value, np.ndarray):
            if value.dtype.kind in "fiub" and value.size <= 256:
                return {"__массив__": list(value.shape), "dtype": str(value.dtype),
                        "значения": [convert(v, path, depth + 1, seen) for v in value.ravel().tolist()]}
            return {"__массив__": list(value.shape), "dtype": str(value.dtype),
                    "sha256": sha256(np.ascontiguousarray(value).tobytes()) if value.dtype.kind != "O" else None}
        if isinstance(value, enum.Enum):
            return f"{type(value).__qualname__}.{value.name}"
        if isinstance(value, (dt.datetime, dt.date, dt.time)):
            return value.isoformat()
        if isinstance(value, Path):
            return str(value)
        if depth > 10 or id(value) in seen:
            return f"<{type(value).__qualname__}: глубже не идём>"
        seen = seen | {id(value)}
        if isinstance(value, dict):
            return {str(k): convert(v, f"{path}.{k}", depth + 1, seen) for k, v in value.items()}
        if isinstance(value, (list, tuple)):
            return [convert(v, f"{path}.{i}", depth + 1, seen) for i, v in enumerate(value)]
        if isinstance(value, (set, frozenset)):
            return sorted(repr(v) for v in value)
        if dataclasses.is_dataclass(value) and not isinstance(value, type):
            body = {"__класс__": type(value).__qualname__}
            for field in dataclasses.fields(value):
                body[field.name] = convert(getattr(value, field.name), f"{path}.{field.name}", depth + 1, seen)
            return body
        module = type(value).__module__ or ""
        if module.startswith("thermogar") and hasattr(value, "__dict__"):
            body = {"__класс__": type(value).__qualname__}
            for name, item in vars(value).items():
                body[name] = convert(item, f"{path}.{name}", depth + 1, seen)
            return body
        return f"<{module}.{type(value).__qualname__}>"

    state = at.session_state.filtered_state
    body = {}
    for key in sorted(state, key=str):
        body[str(key)] = convert(state[key], str(key), 0, set())
    files = {f"{prefix}.json": json.dumps(body, ensure_ascii=False, indent=1).encode("utf-8")}
    files.update(tables)
    return files


# ---------------------------------------------------------------------------
# Драйвер приложения
# ---------------------------------------------------------------------------


class App:
    def __init__(self, driver: "Driver", name: str, at: Any, env: dict[str, str | None]) -> None:
        self.driver = driver
        self.name = name
        self.at = at
        self.env = env

    @property
    def state(self) -> Any:
        return self.at.session_state

    def run(self, action: str) -> "App":
        self.driver.run(self, action)
        return self

    def click(self, key: str, action: str | None = None) -> "App":
        self.at.button(key=key).click()
        return self.run(action or f"нажать {key}")


class Driver:
    def __init__(self, recorder: Recorder, state_base: Path) -> None:
        self.recorder = recorder
        self.state_base = state_base
        self.apps: dict[str, App] = {}

    def new(self, name: str = "app", *, state: dict[str, Any] | None = None,
            state_dir: str = "state", script: Path = APP_PATH,
            local_app_data: str | None = None) -> App:
        from streamlit.testing.v1 import AppTest

        if local_app_data is None:
            env = {"THERMOGAR_STATE_ROOT": str(self.state_base / state_dir), "LOCALAPPDATA": None}
        else:
            env = {"THERMOGAR_STATE_ROOT": "", "LOCALAPPDATA": str(self.state_base / local_app_data)}
        self._apply(env)
        at = AppTest.from_file(str(script), default_timeout=APP_TIMEOUT_S)
        for key, value in (state or {}).items():
            at.session_state[key] = value
        app = App(self, name, at, env)
        self.apps[name] = app
        return app

    def _apply(self, env: dict[str, str | None]) -> None:
        for key, value in env.items():
            if value is None:
                continue
            if value == "":
                os.environ.pop(key, None)
            else:
                os.environ[key] = value

    def run(self, app: App, action: str) -> None:
        self._apply(app.env)
        started = time.perf_counter()
        app.at.run()
        self.recorder.snapshot(app.name, app.at, action, time.perf_counter() - started)


# ---------------------------------------------------------------------------
# Помощники сценариев
# ---------------------------------------------------------------------------


def _import_tools(name: str) -> Any:
    if str(TOOLS_DIR) not in sys.path:
        sys.path.insert(0, str(TOOLS_DIR))
    return __import__(name)


def sidebar_state(key: str, composition: str, units: str, balance: str) -> dict[str, Any]:
    return {
        "thermogar_database_key": key,
        f"thermogar_composition_{key}": composition,
        f"thermogar_units_{key}": units,
        f"thermogar_balance_{key}": balance,
    }


def f_start(d: Driver, key: str, session: dict[str, Any] | None = None,
            prepare: Callable[[App], None] | None = None) -> App:
    """Как ``start`` в tools/test_ui_f.py: профиль BASES, затем подготовка."""

    f = _import_tools("test_ui_f")
    profile = f.BASES[key]
    state = sidebar_state(key, profile["composition"], profile["units"], profile["balance"])
    state.update(session or {})
    app = d.new(state=state).run("запуск")
    if prepare is not None:
        prepare(app)
        app.run("подготовка")
    return app


def fill_editor(values_for: Callable[[Any], dict[str, Any]]) -> None:
    """Подмена st.data_editor, как в tools/test_ui_f.py (AppTest не вводит в редактор)."""

    import streamlit as st

    original = st.data_editor

    def filled(frame: Any, *args: Any, **kwargs: Any) -> Any:
        original(frame, *args, **kwargs)
        edited = frame.copy()
        for column, value in values_for(edited).items():
            edited[column] = value
        return edited

    st.data_editor = filled


# -- А. tools/test_ui_f.py ----------------------------------------------------


def case_f_startup(d: Driver, key: str) -> None:
    f_start(d, key)


def case_f_single(d: Driver, key: str) -> None:
    f = _import_tools("test_ui_f")
    f_start(d, key, {f"single_temperature_{key}": f.BASES[key]["temperature"]}).click("single_calculate")


def case_f_tscan(d: Driver, key: str) -> None:
    f = _import_tools("test_ui_f")
    t_min, t_max, t_step = f.BASES[key]["scan"]
    f_start(d, key, {f"t_min_{key}": t_min, f"t_max_{key}": t_max, f"t_step_{key}": t_step}).click(
        "temperature_calculate")


def case_f_cscan(d: Driver, key: str) -> None:
    f = _import_tools("test_ui_f")
    c_min, c_max, c_step = f.BASES[key]["concentration"]
    variable = f.BASES[key]["variable"]
    app = f_start(d, key, prepare=lambda a: f.set_select(a.at, "Изменяемый элемент", variable))
    # 21-Ж: символ элемента и единицы — в подписи поля.
    symbol = variable[:1] + variable[1:].lower()
    suffix = "ат.%" if f.BASES[key]["units"] == "атомные %" else "мас.%"
    f.set_number(app.at, f"{symbol}: от, {suffix}", c_min)
    f.set_number(app.at, f"{symbol}: до, {suffix}", c_max)
    f.set_number(app.at, f"{symbol}: шаг, {suffix}", c_step)
    app.run("задать диапазон состава")
    app.click("concentration_calculate")


def case_f_solid(d: Driver, key: str, method: str) -> None:
    f = _import_tools("test_ui_f")
    profile = f.BASES[key]
    start_key = (f"solidification_start_13_1_{key}_thermogar_patch" if key == "fe"
                 else f"solidification_start_13_1_{key}")
    f_start(d, key, {
        f"solidification_method_{key}": method,
        start_key: profile["solidification_start"],
        f"solidification_step_{key}": profile["solidification_step"],
    }).click("solidification_calculate")


def case_f_energy(d: Driver, key: str) -> None:
    f_start(d, key, {f"energy_t_min_{key}": 600.0, f"energy_t_max_{key}": 1000.0,
                     f"energy_t_step_{key}": 100.0}).click("energy_curve_calculate")


def case_f_driving(d: Driver, key: str) -> None:
    f_start(d, key, {f"driving_t_min_{key}": 600.0, f"driving_t_max_{key}": 1000.0,
                     f"driving_t_step_{key}": 100.0}).click("driving_force_calculate")


def case_f_tzero(d: Driver, key: str) -> None:
    f = _import_tools("test_ui_f")
    t_min, t_max = f.TZERO_WINDOWS[key]
    f_start(d, key, {f"tzero_t_min_{key}": t_min, f"tzero_t_max_{key}": t_max}).click("tzero_calculate")


def case_f_density(d: Driver, key: str) -> None:
    f = _import_tools("test_ui_f")
    f_start(d, key, {f"physical_temperature_{key}": f.BASES[key]["temperature"]}).click(
        "physical_single_calculate")


def case_f_density_warning(d: Driver) -> None:
    f = _import_tools("test_ui_f")
    f_start(d, "ni", {
        "thermogar_composition_ni": ("C=0.005, SI=0.10, MN=0.50, S=0.020, CR=23.5, MO=13.0, "
                                     "NB=0.06, AL=0.25, TI=0.10, FE=0.50"),
        "thermogar_units_ni": "массовые %",
        "thermogar_balance_ni": "NI",
        "physical_temperature_ni": 25.0,
        "physical_single_manual": True,
        "physical_single_tokens": f._phases_without_bcc_b2("ni"),
    }).click("physical_single_calculate")


def case_f_density_t(d: Driver, key: str) -> None:
    f = _import_tools("test_ui_f")
    t_min, _t_max, t_step = f.BASES[key]["scan"]
    f_start(d, key, {f"physical_t_min_{key}": t_min, f"physical_t_max_{key}": t_min + 2.0 * t_step,
                     f"physical_t_step_{key}": t_step}).click("physical_scan_calculate")


def case_f_elastic(d: Driver, key: str) -> None:
    f = _import_tools("test_ui_f")
    fill_editor(lambda _frame: dict(f.ELASTIC_ROW_VALUES))
    app = f_start(d, key, {f"b4b2_elastic_temperature_{key}": f.BASES[key]["temperature"]})
    app.click("b4b2_elastic_prepare_calculate")
    app.click("b4b2_elastic_vrh_calculate")


def case_f_strength(d: Driver, key: str) -> None:
    f_start(d, key, {
        f"b4b2_strengthening_rule_{key}": "Линейная сумма",
        f"b4b2_strengthening_confirmation_{key}": True,
        f"b4b2_strengthening_provenance_{key}": "учебные значения",
        f"b4b2_hall_use_{key}": True,
        f"b4b2_taylor_use_{key}": True,
    }).click("b4b2_strengthening_calculate")


def case_f_click(d: Driver, key: str, button: str) -> None:
    f_start(d, key).click(button)


def case_f_db_change(d: Driver) -> None:
    f = _import_tools("test_ui_f")
    app = f_start(d, "ni", {"single_temperature_ni": f.BASES["ni"]["temperature"]})
    app.click("single_calculate")
    app.state["thermogar_database_key"] = "fe"
    app.state["thermogar_composition_fe"] = f.BASES["fe"]["composition"]
    app.state["thermogar_units_fe"] = f.BASES["fe"]["units"]
    app.run("сменить базу на fe")


# -- Б. tools/test_ui_g.py ----------------------------------------------------


def g_start(d: Driver, key: str, composition: str | None = None, name: str = "app") -> App:
    g = _import_tools("test_ui_g")
    balance, default_composition, units = g.SIDEBAR_ALLOY[key]
    state = sidebar_state(key, default_composition if composition is None else composition, units, balance)
    return d.new(name, state=state).run("запуск")


def g_diffusion_inputs(app: App, key: str, prefix: str) -> None:
    g = _import_tools("test_ui_g")
    balance, units, left, right, temperature, _phases = g.DIFFUSION_COUPLE[key]
    state = app.state
    state[f"{prefix}_balance_{key}"] = balance
    state[f"{prefix}_units_{key}"] = units
    state[f"{prefix}_left_{key}"] = left
    state[f"{prefix}_right_{key}"] = right
    state[f"{prefix}_temperature_{key}"] = temperature
    state[f"{prefix}_length_{key}"] = 100.0
    state[f"{prefix}_time_{key}"] = g.SHORT_TIME_H
    state[f"{prefix}_nodes_{key}"] = 20
    app.run(f"входы пары {prefix}")


def case_g_render(d: Driver, key: str) -> None:
    g_start(d, key)


def case_g_run_gate(d: Driver) -> None:
    for key in ("ni", "al", "fe"):
        g_start(d, key, name=f"app_{key}")


def case_g_single(d: Driver, key: str) -> None:
    g = _import_tools("test_ui_g")
    app = g_start(d, key)
    g_diffusion_inputs(app, key, "kin_single")
    app.state[f"kin_single_phase_{key}"] = g.DIFFUSION_COUPLE[key][5][0]
    app.run("выбрать фазу")
    app.click(f"kin_single_run_{key}")


def case_g_hom(d: Driver, key: str) -> None:
    g = _import_tools("test_ui_g")
    app = g_start(d, key)
    g_diffusion_inputs(app, key, "kin_hom")
    app.state[f"kin_hom_phases_{key}"] = list(g.DIFFUSION_COUPLE[key][5][:2])
    app.run("выбрать фазы")
    app.click(f"kin_hom_run_{key}")


def case_g_hom_al(d: Driver) -> None:
    app = g_start(d, "al")
    g_diffusion_inputs(app, "al", "kin_hom")


def g_kwn_state(app: App, key: str, matrix: str, precipitate: str, temperature: float | None,
                duration_h: float, bins: int | None) -> None:
    state = app.state
    state[f"precipitation_{key}_user_matrix"] = matrix
    state[f"precipitation_{key}_user_precipitate"] = precipitate
    if temperature is not None:
        state[f"precipitation_{key}_user_temperature_c"] = temperature
    state[f"precipitation_{key}_user_duration_h"] = duration_h
    if bins is not None:
        state[f"precipitation_{key}_user_bins"] = bins


def case_g_kwn(d: Driver, key: str) -> None:
    g = _import_tools("test_ui_g")
    matrix, precipitate, temperature = g.KWN_CELL[key]
    app = g_start(d, key)
    g_kwn_state(app, key, matrix, precipitate, temperature, g.SHORT_TIME_H, g.KWN_BINS)
    app.run("входы KWN")
    app.click(f"precipitation_{key}_user_calculate")


def case_g_fe_provenance(d: Driver) -> None:
    g = _import_tools("test_ui_g")
    matrix, precipitate, temperature = g.KWN_CELL["fe"]
    app = g_start(d, "fe")
    g_kwn_state(app, "fe", matrix, precipitate, temperature, g.SHORT_TIME_H, 30)
    app.run("входы KWN, 30 классов")
    app.click("precipitation_fe_user_calculate")


def case_g_ni_kwn_hour(d: Driver) -> None:
    g = _import_tools("test_ui_g")
    matrix, precipitate, temperature = g.KWN_CELL["ni"]
    app = g_start(d, "ni")
    g_kwn_state(app, "ni", matrix, precipitate, temperature, 1.0, None)
    app.run("входы KWN, 1 ч")
    app.click("precipitation_ni_user_calculate")


def case_g_too_long(d: Driver) -> None:
    g = _import_tools("test_ui_g")
    g_start(d, "fe", composition=g.FE_EIGHT_SOLUTES + ", NB=0.05, TI=0.02, AL=0.02")


def case_g_eight(d: Driver) -> None:
    g = _import_tools("test_ui_g")
    g_start(d, "fe", composition=g.FE_EIGHT_SOLUTES)


def case_g_bad_couple(d: Driver, left: str, right: str) -> None:
    app = g_start(d, "ni")
    g_diffusion_inputs(app, "ni", "kin_single")
    app.state["kin_single_left_ni"] = left
    app.state["kin_single_right_ni"] = right
    app.run("неверная пара")


def case_g_kwn_grid(d: Driver) -> None:
    g = _import_tools("test_ui_g")
    app = g_start(d, "fe")
    state = app.state
    state["precipitation_fe_user_matrix"] = "BCC_A2"
    state["precipitation_fe_user_precipitate"] = "M23C6"
    state["precipitation_fe_user_duration_h"] = g.SHORT_TIME_H
    state["precipitation_fe_user_cmin_nm"] = 10.0
    state["precipitation_fe_user_cmax_nm"] = 1.0
    app.run("входы KWN, радиусы наоборот")
    app.click("precipitation_fe_user_calculate")


# -- В. tools/test_ui_h.py и «Проекты и данные» -------------------------------


def h_check_upstream() -> Any:
    """Сценарии test_ui_h идут по самому ThermoGar_app.py, если его образцы не совпали."""

    h = _import_tools("test_ui_h")
    source = APP_PATH.read_text(encoding="utf-8")
    for broken, _fixed in h.UPSTREAM_PATCHES:
        if broken in source:
            raise SystemExit("образец UPSTREAM_PATCHES из test_ui_h совпал: сценарий шёл бы по копии")
    return h


def h_start(d: Driver, name: str = "app", state_dir: str = "state") -> App:
    h_check_upstream()
    return d.new(name, state_dir=state_dir).run("запуск")


def h_labelled(elements: Any, label: str) -> Any:
    return _import_tools("test_ui_h").labelled(elements, label)


def h_widget(elements: Any, key: str) -> Any:
    return _import_tools("test_ui_h").widget(elements, key)


def case_h_library(d: Driver, key: str) -> None:
    h = h_check_upstream()
    app = h_start(d)
    at = app.at
    at.sidebar.selectbox("thermogar_database_key").set_value(key)
    app.run(f"база {key}")
    _balance, units, _t, composition = h.CASES[key]
    app.state[f"thermogar_composition_{key}"] = composition
    app.state[f"thermogar_units_{key}"] = units.replace("ат.%", "атомные %").replace("мас.%", "массовые %")
    app.run("состав")
    h_labelled(at.text_input, "Название марки или состава").set_value(f"Проверка {key}")
    h_labelled(at.button, "Сохранить текущий состав").click()
    app.run("сохранить состав")
    saved = json.loads((d.state_base / "state" / "workspace" / "alloys.json").read_text(encoding="utf-8"))["alloys"]
    other = "ni" if key != "ni" else "al"
    at.sidebar.selectbox("thermogar_database_key").set_value(other)
    app.run(f"база {other}")
    h_widget(at.selectbox, "alloy_selected_id").set_value(saved[0]["id"])
    app.run("выбрать запись")
    h_widget(at.button, "alloy_load_button").click()
    app.run("загрузить запись")
    h_widget(at.checkbox, f"alloy_delete_confirm_{saved[0]['id']}").set_value(True)
    app.run("подтвердить удаление")
    h_widget(at.button, f"alloy_delete_button_{saved[0]['id']}").click()
    app.run("удалить запись")


def newest_artifact(root: Path, kind: str) -> bytes:
    newest: tuple[int, Path] | None = None
    for path in (root / "state").rglob("*"):
        if path.is_file() and path.parent.name == kind:
            stamp = path.stat().st_mtime_ns
            if newest is None or stamp >= newest[0]:
                newest = (stamp, path)
    assert newest is not None, kind
    return newest[1].read_bytes()


def case_h_library_roundtrip(d: Driver) -> None:
    app = h_start(d)
    h_labelled(app.at.text_input, "Название марки или состава").set_value("Сталь")
    h_labelled(app.at.button, "Сохранить текущий состав").click()
    app.run("сохранить состав")
    exported = newest_artifact(d.state_base / "state", "alloy-library-json-v1")
    second = h_start(d, "app2", "state2")
    h_labelled(second.at.file_uploader, "Импортировать библиотеку JSON").set_value(
        ("alloys.json", exported, "application/json"))
    second.run("загрузить файл библиотеки")
    h_widget(second.at.button, "alloy_import_button").click()
    second.run("импорт библиотеки")


def case_h_project(d: Driver, key: str) -> None:
    app = h_start(d)
    app.at.sidebar.selectbox("thermogar_database_key").set_value(key)
    app.run(f"база {key}")
    app.at.number_input(f"single_temperature_{key}").set_value(812.0)
    app.run("температура 812")
    h_labelled(app.at.text_input, "Название проекта").set_value("Проект H")
    h_labelled(app.at.button, "Сохранить проект в папке ThermoGar").click()
    app.run("сохранить проект")
    fresh = h_start(d, "app2", "state")
    h_widget(fresh.at.button, "project_load_button").click()
    fresh.run("открыть проект")


def case_h_project_export(d: Driver) -> None:
    app = h_start(d)
    app.at.number_input("single_temperature_fe").set_value(755.0)
    app.run("температура 755")
    h_labelled(app.at.text_input, "Название проекта").set_value("Проект H")
    h_labelled(app.at.button, "Сохранить проект в папке ThermoGar").click()
    app.run("сохранить проект")
    exported = newest_artifact(d.state_base / "state", "project-json-v1")
    second = h_start(d, "app2", "state2")
    h_labelled(second.at.file_uploader, "Импортировать проект").set_value(
        ("project.json", exported, "application/json"))
    second.run("загрузить файл проекта")
    h_widget(second.at.button, "project_import_button").click()
    second.run("импорт проекта")
    h_widget(second.at.checkbox, "project_delete_confirm").set_value(True)
    second.run("подтвердить удаление")
    h_widget(second.at.button, "project_delete_button").click()
    second.run("удалить проект")


def case_h_confirmations(d: Driver) -> None:
    app = h_start(d)
    h_labelled(app.at.text_input, "Название проекта").set_value("Проект H")
    h_labelled(app.at.button, "Сохранить проект в папке ThermoGar").click()
    app.run("сохранить проект")
    h_labelled(app.at.text_input, "Название марки или состава").set_value("Марка H")
    h_labelled(app.at.button, "Сохранить текущий состав").click()
    app.run("сохранить состав")


def case_h_history(d: Driver) -> None:
    app = h_start(d)
    at = app.at
    h_labelled(at.text_input, "Название марки или состава").set_value("Сталь")
    h_labelled(at.button, "Сохранить текущий состав").click()
    app.run("сохранить состав")
    h_widget(at.radio, "projects_history_mode").set_value("История расчётов")
    app.run("история расчётов")
    h_widget(at.button, "history_restore_button").click()
    app.run("восстановить из истории")
    h_widget(at.radio, "projects_history_mode").set_value("История расчётов")
    app.run("история расчётов")
    h_widget(at.checkbox, "history_clear_confirm").set_value(True)
    app.run("подтвердить очистку")
    h_widget(at.button, "history_clear_button").click()
    app.run("очистить историю")


def h_upload_batch(app: App, payload: bytes) -> None:
    h_labelled(app.at.file_uploader, "Файл составов").set_value(("batch.csv", payload, "text/csv"))
    app.run("загрузить файл составов")


def case_h_batch_sep(d: Driver, separator: str, bom: bool) -> None:
    h = h_check_upstream()
    h_upload_batch(h_start(d), h.batch_csv(separator=separator, bom=bom))


def case_h_batch_junk(d: Driver) -> None:
    h_upload_batch(h_start(d), b"foo,bar\n1,2\n")


def case_h_batch_template(d: Driver) -> None:
    app = h_start(d)
    app.click("batch_template_xlsx_prepare")
    app.click("batch_template_csv_prepare")


def case_h_batch_summary(d: Driver) -> None:
    h = h_check_upstream()
    app = h_start(d)
    h_upload_batch(app, h.batch_csv(("ni",)))
    app.click("batch_calculate_button")


def case_h_batch_three(d: Driver) -> None:
    h = h_check_upstream()
    app = h_start(d)
    h_upload_batch(app, h.batch_csv())
    app.click("batch_calculate_button")
    app.click("batch_result_export_prepare")


def case_h_batch_c15(d: Driver) -> None:
    payload = ('Название,База,Основа,Единицы,"Температура, °C",Добавки,Фазы\n'
               'FE,fe,FE,мас.%,700,"C=0.2, CR=11.5",C15_LAVES\n').encode("utf-8")
    h_upload_batch(h_start(d), payload)


def case_h_first_run(d: Driver) -> None:
    h_check_upstream()
    d.new(local_app_data="LocalAppData").run("запуск без THERMOGAR_STATE_ROOT")


def case_h_phase_reference(d: Driver, key: str) -> None:
    app = d.new(state={"thermogar_database_key": key}).run("запуск")
    h_labelled(app.at.checkbox, "Показывать исходное английское описание").set_value(True)
    app.run("показать английское описание")


def case_h_batch_fe08(d: Driver) -> None:
    """Пакет Fe–0,8C мас. %, 700 °C, 4 строки — входы (б) приёмки 19-Д, через экран."""

    import thermogar_verified_state as vs

    rows = [("Fe-пусто", ""), ("Fe-метастабильный", "метастабильный"),
            ("Fe-стабильный", "стабильный"), ("Fe-metastabe", "metastabe")]
    out = io.StringIO(newline="")
    writer = csv.writer(out, delimiter=",", lineterminator="\n")
    writer.writerow(vs.TEMPLATE_HEADERS)
    for name, mode in rows:
        writer.writerow((name, "fe", "FE", "мас.%", 700, "C=0.8", mode, 101325, ""))
    app = h_start(d)
    h_upload_batch(app, out.getvalue().encode("utf-8"))
    if any(button.key == "batch_calculate_button" for button in app.at.button):
        app.click("batch_calculate_button")
        if any(button.key == "batch_result_export_prepare" for button in app.at.button):
            app.click("batch_result_export_prepare")


# -- Г. составы 15-Ш -----------------------------------------------------------

SH_CASES = {
    # results/wave15_sh/scripts/density_run.py: база, состав, единицы, основа, T, °C
    "nicr": ("ni", "CR=20", "массовые %", "NI", 700.0),
    "nialcr": ("ni", "AL=9.8, CR=8.3", "атомные %", "NI", 800.0),
    "fecrc": ("fe", "CR=15, C=0.4", "массовые %", "FE", 950.0),
}
SH_SCAN = (100.0, 1100.0, 100.0)
SH_YOUNG_BY_ORDER = (100.0, 300.0)


def sh_start(d: Driver, case: str, overrides: str, extra: dict[str, Any]) -> App:
    key, composition, units, balance, _t = SH_CASES[case]
    state = sidebar_state(key, composition, units, balance)
    if overrides == "off":
        state["physical_overrides_enabled"] = False
    state.update(extra)
    return d.new(state=state).run("запуск")


def case_sh_density(d: Driver, case: str, overrides: str) -> None:
    key, *_rest, temperature = SH_CASES[case]
    sh_start(d, case, overrides, {f"physical_temperature_{key}": temperature}).click("physical_single_calculate")


def case_sh_density_t(d: Driver, case: str, overrides: str) -> None:
    key = SH_CASES[case][0]
    sh_start(d, case, overrides, {f"physical_t_min_{key}": SH_SCAN[0], f"physical_t_max_{key}": SH_SCAN[1],
                                  f"physical_t_step_{key}": SH_SCAN[2]}).click("physical_scan_calculate")


def case_sh_elastic(d: Driver, case: str, overrides: str) -> None:
    key, *_rest, temperature = SH_CASES[case]

    def values(frame: Any) -> dict[str, Any]:
        # results/wave15_f/scripts/elastic_run.py: условные модули 100/300 ГПа, ν = 0,25.
        return {"young_gpa": [SH_YOUNG_BY_ORDER[i % 2] for i in range(len(frame))], "poisson": 0.25,
                "origin": "measured", "source": "условное значение для проверки весов VRH",
                "reference_temperature_c": 25.0}

    fill_editor(values)
    app = sh_start(d, case, overrides, {f"b4b2_elastic_temperature_{key}": temperature})
    app.click("b4b2_elastic_prepare_calculate")
    app.click("b4b2_elastic_vrh_calculate")


def case_sh_solid(d: Driver, case: str) -> None:
    # results/wave18_v/run_case.py: умолчания вкладки, pdens 50.
    key = SH_CASES[case][0]
    sh_start(d, case, "on", {f"solidification_pdens_{key}": 50}).click("solidification_calculate")


# ---------------------------------------------------------------------------
# Перечень случаев
# ---------------------------------------------------------------------------

DBS = ("ni", "al", "fe")
SOLID_METHODS = {
    "sravn": "Сравнить равновесное и Scheil–Gulliver",
    "ravn": "Только равновесное затвердевание",
    "scheil": "Только Scheil–Gulliver",
}


@dataclasses.dataclass(frozen=True)
class Case:
    name: str
    group: str
    source: str
    run: Callable[[Driver], None]
    slow: bool = False


def build_cases() -> list[Case]:
    cases: list[Case] = []

    def add(name: str, group: str, source: str, fn: Callable[..., None], *args: Any, slow: bool = False) -> None:
        cases.append(Case(name, group, source, lambda d, fn=fn, args=args: fn(d, *args), slow))

    f = "tools/test_ui_f.py::"
    for key in DBS:
        add(f"f_startup_{key}", "А", f + f"test_startup_is_clean[{key}]", case_f_startup, key)
    for key in DBS:
        add(f"f_single_{key}", "А", f + f"test_single_equilibrium[{key}]", case_f_single, key)
    for key in DBS:
        add(f"f_tscan_{key}", "А", f + f"test_temperature_scan[{key}]", case_f_tscan, key)
    for key in DBS:
        add(f"f_cscan_{key}", "А", f + f"test_concentration_scan[{key}]", case_f_cscan, key)
    for key in DBS:
        for tag, method in SOLID_METHODS.items():
            add(f"f_solid_{key}_{tag}", "А", f + f"test_solidification[{key}-{method}]", case_f_solid, key, method,
                slow=True)
    for key in DBS:
        add(f"f_energy_{key}", "А", f + f"test_energy_curve[{key}]", case_f_energy, key)
    for key in DBS:
        add(f"f_driving_{key}", "А", f + f"test_driving_force[{key}]", case_f_driving, key)
    for key in DBS:
        add(f"f_tzero_{key}", "А", f + f"test_tzero_in_narrow_window[{key}]", case_f_tzero, key)
    for key in DBS:
        add(f"f_density_{key}", "А", f + f"test_density_single[{key}]", case_f_density, key)
    add("f_density_warning_ni", "А", f + "test_density_estimated_warning_is_shown_to_user", case_f_density_warning)
    for key in DBS:
        add(f"f_density_t_{key}", "А", f + f"test_density_temperature_scan[{key}]", case_f_density_t, key)
    for key in DBS:
        add(f"f_elastic_{key}", "А", f + f"test_elastic_vrh[{key}]", case_f_elastic, key)
    for key in DBS:
        add(f"f_strength_{key}", "А", f + f"test_strengthening[{key}]", case_f_strength, key)
    for test, button, tag in (("test_binary_diagram", "binary_calculate", "binary"),
                              ("test_isopleth_diagram", "isopleth_calculate", "isopleth"),
                              ("test_ternary_diagram", "ternary_calculate", "ternary"),
                              ("test_ternary_phase_map", "ternary_map_calculate", "tmap")):
        for key in DBS:
            add(f"f_{tag}_{key}", "А", f + f"{test}[{key}]", case_f_click, key, button, slow=True)
    for key in DBS:
        add(f"f_phasemap3_{key}", "А", f + f"test_phase_map_needs_three_elements[{key}]", case_f_startup, key)
    add("f_db_change", "А", f + "test_results_do_not_survive_a_database_change", case_f_db_change)

    g = "tools/test_ui_g.py::"
    for key in DBS:
        add(f"g_render_{key}", "Б", g + f"test_kinetics_section_renders[{key}]", case_g_render, key)
    add("g_run_gate", "Б", g + "test_run_gate_is_identical_for_all_databases", case_g_run_gate)
    for key in DBS:
        add(f"g_single_{key}", "Б", g + f"test_diffusion_single_phase[{key}]", case_g_single, key)
    for key in ("ni", "fe"):
        add(f"g_hom_{key}", "Б", g + f"test_diffusion_homogenization[{key}]", case_g_hom, key)
    add("g_hom_al", "Б", g + "test_homogenization_unavailable_on_al_is_explained", case_g_hom_al)
    for key in DBS:
        add(f"g_kwn_{key}", "Б", g + f"test_kwn_precipitation[{key}]", case_g_kwn, key, slow=key == "fe")
    add("g_kwn_fe_provenance", "Б", g + "test_fe_kwn_provenance_status_is_neutral", case_g_fe_provenance, slow=True)
    for key in ("al", "fe"):
        add(f"g_kwn_matrix_{key}", "Б", g + f"test_kwn_matrix_offers_the_disordered_half[{key}]", case_g_render, key)
    add("g_kwn_too_long", "Б", g + "test_kwn_reports_a_too_long_composition_without_a_traceback", case_g_too_long)
    add("g_kwn_eight", "Б", g + "test_kwn_accepts_the_eight_solute_steel_since_bl21", case_g_eight)
    for prefix in ("kin_single", "kin_hom"):
        add(f"g_bounded_{prefix}", "Б", g + f"test_diffusion_number_inputs_are_bounded[{prefix}]", case_g_render, "ni")
    for number, (left, right) in enumerate((("CR=7.7", "CR=7.7"), ("CR=60, AL=50", "CR=1"),
                                            ("бессмыслица", "CR=5")), start=1):
        add(f"g_bad_couple_{number}", "Б", g + "test_bad_couple_composition_is_reported", case_g_bad_couple,
            left, right)
    add("g_kwn_grid", "Б", g + "test_kwn_size_grid_is_validated", case_g_kwn_grid)
    add("g_ni_kwn_hour", "Б", g + "test_ni_kwn_reaches_a_precipitated_state", case_g_ni_kwn_hour, slow=True)
    add("g_fe_defaults", "Б", g + "test_fe_shipped_defaults_are_the_declared_ones", case_g_render, "fe")

    h = "tools/test_ui_h.py::"
    for key in DBS:
        add(f"h_library_{key}", "В", h + f"test_library_save_appears_in_the_list_and_loads_back[{key}]",
            case_h_library, key)
    add("h_library_roundtrip", "В", h + "test_library_export_and_import_round_trip", case_h_library_roundtrip)
    for key in DBS:
        add(f"h_project_{key}", "В", h + f"test_project_saves_fourteen_keys_and_restores_state[{key}]",
            case_h_project, key)
    add("h_project_export", "В", h + "test_project_export_is_portable_and_imports_back", case_h_project_export)
    add("h_confirmations", "В", h + "test_confirmations_survive_the_rerun_and_stay_in_their_own_section",
        case_h_confirmations)
    add("h_history", "В", h + "test_history_records_events_exports_csv_and_clears", case_h_history)
    for tag, separator, bom in (("comma", ",", False), ("semicolon", ";", False),
                                ("comma_bom", ",", True), ("semicolon_bom", ";", True)):
        add(f"h_batch_{tag}", "В", h + f"test_batch_accepts_both_separators_and_both_encodings[{tag}]",
            case_h_batch_sep, separator, bom)
    add("h_batch_junk", "В", h + "test_batch_rejects_a_junk_file_with_a_readable_message", case_h_batch_junk)
    add("h_batch_template", "В", h + "test_batch_template_downloads_open_as_excel_and_csv", case_h_batch_template)
    add("h_batch_summary", "В", h + "test_batch_summary_has_no_receipt_columns", case_h_batch_summary)
    add("h_batch_three", "В", h + "test_batch_calculates_three_databases_and_exports_excel", case_h_batch_three,
        slow=True)
    add("h_batch_c15", "В", h + "test_batch_rejects_c15_for_steel_before_any_calculation", case_h_batch_c15,
        slow=True)
    add("h_first_run", "В", h + "test_first_run_creates_the_profile_and_writes_nothing_into_the_program",
        case_h_first_run)
    for key in DBS:
        add(f"h_phase_reference_{key}", "В", "справочник фаз", case_h_phase_reference, key)
    add("h_batch_fe08", "В", "results/wave19_d/priemka.py::case_batch (через экран)", case_h_batch_fe08)

    sh = "results/wave15_sh/scripts/"
    for case in SH_CASES:
        for overrides in ("on", "off"):
            add(f"sh_density_{case}_{overrides}", "Г", sh + "density_run.py", case_sh_density, case, overrides)
            add(f"sh_density_t_{case}_{overrides}", "Г", sh + "density_run.py (скан)", case_sh_density_t, case,
                overrides)
            add(f"sh_elastic_{case}_{overrides}", "Г", "results/wave15_f/scripts/elastic_run.py", case_sh_elastic,
                case, overrides)
    for case in SH_CASES:
        add(f"sh_solid_{case}", "Г", f"results/wave18_v/run_case.py::{case}", case_sh_solid, case, slow=True)
    return cases


# ---------------------------------------------------------------------------
# Один случай в своём процессе
# ---------------------------------------------------------------------------


def package_versions() -> dict[str, str]:
    versions = {"python": sys.version.split()[0]}
    for name in PACKAGES:
        try:
            versions[name] = importlib.metadata.version(name)
        except importlib.metadata.PackageNotFoundError:
            versions[name] = "нет"
    return versions


def run_case(name: str, out: Path, state_base: Path) -> int:
    cases = {case.name: case for case in build_cases()}
    case = cases[name]
    if out.exists():
        raise SystemExit(f"каталог уже есть, не перезаписываю: {out}")
    for variable, expected in (("PYTHONHASHSEED", "0"), ("MPLBACKEND", "Agg")):
        if os.environ.get(variable) != expected:
            raise SystemExit(f"нужно {variable}={expected}")
    os.environ.pop("THERMOGAR_PHYSICAL_OVERRIDES", None)
    state_base.mkdir(parents=True, exist_ok=True)
    os.chdir(PROJECT_ROOT)
    if str(APP_DIR) not in sys.path:
        sys.path.insert(0, str(APP_DIR))

    import matplotlib

    matplotlib.use("Agg")
    import streamlit as st

    recorder = Recorder(out)
    original_download = st.download_button

    def capture(label: Any, data: Any = None, *args: Any, **kwargs: Any) -> Any:
        recorder.capture_download(str(kwargs.get("file_name") or label), data)
        return original_download(label, data, *args, **kwargs)

    st.download_button = capture

    # Пул — как фикстура _bounded_memory в tools/test_ui_f.py.
    f = _import_tools("test_ui_f")
    import thermogar_parallel_ui

    thermogar_parallel_ui._WORKER_COUNT = f.UI_TEST_POOL_WORKERS

    driver = Driver(recorder, state_base)
    started_wall = dt.datetime.now().astimezone()
    started = time.perf_counter()
    error = None
    try:
        case.run(driver)
    except BaseException as caught:  # noqa: BLE001
        error = f"{type(caught).__name__}: {caught}"
    finally:
        import gc

        import matplotlib.pyplot as plt

        thermogar_parallel_ui.close_shared_engines()
        plt.close("all")
        gc.collect()
    app_bytes = APP_PATH.read_bytes()
    meta = {
        "случай": case.name,
        "группа": case.group,
        "источник": case.source,
        "slow": case.slow,
        "ошибка_сценария": error,
        "воркеров_пула": f.UI_TEST_POOL_WORKERS,
        "ThermoGar_app.py_sha256": sha256(app_bytes),
        "версии": package_versions(),
        "начало": started_wall.isoformat(timespec="seconds"),
        "конец": dt.datetime.now().astimezone().isoformat(timespec="seconds"),
        "секунд": round(time.perf_counter() - started, 1),
        "папка_состояния": str(state_base),
    }
    recorder.write({name: app.at for name, app in driver.apps.items()}, meta)
    print(json.dumps({"случай": case.name, "секунд": meta["секунд"], "ошибка": error,
                      "исключений_на_экране": recorder.exceptions}, ensure_ascii=False))
    return 1 if error else 0


# ---------------------------------------------------------------------------
# Прогон: случаи по одному, память и время
# ---------------------------------------------------------------------------


def run_all(out: Path, state: Path, time_csv: Path, only: str | None, min_free: float,
            wait_s: float) -> int:
    import psutil

    cases = [case for case in build_cases() if only is None or re.search(only, case.name)]
    out.mkdir(parents=True, exist_ok=True)
    new_file = not time_csv.exists()
    handle = time_csv.open("a", encoding="utf-8", newline="")
    writer = csv.writer(handle, lineterminator="\n")
    if new_file:
        writer.writerow(["случай", "группа", "slow", "начало", "секунд", "пик_дерева_гиб",
                         "свободно_на_входе_гиб", "мин_свободно_гиб", "ждал_памяти_с", "код_выхода",
                         "снят_по_памяти"])
    env = dict(os.environ)
    env.update({"MPLBACKEND": "Agg", "PYTHONHASHSEED": "0", "PYTHONDONTWRITEBYTECODE": "1",
                "PYTHONIOENCODING": "utf-8"})
    env.pop("THERMOGAR_PHYSICAL_OVERRIDES", None)
    failures = 0
    run_started = time.perf_counter()
    for number, case in enumerate(cases, start=1):
        target = out / case.name
        if (target / "sha256.txt").exists():
            print(f"[{number}/{len(cases)}] {case.name}: уже снят, пропускаю", flush=True)
            continue
        if target.exists():
            print(f"[{number}/{len(cases)}] {case.name}: неполный каталог {target} — СТОП", flush=True)
            return 2
        waited = 0.0
        free = psutil.virtual_memory().available / GIB
        while free < min_free:
            if waited >= wait_s:
                print(f"СТОП: свободно {free:.2f} ГиБ < {min_free} ГиБ дольше {wait_s:.0f} с", flush=True)
                return 3
            time.sleep(10)
            waited += 10
            free = psutil.virtual_memory().available / GIB
        command = [sys.executable, "-B", "-X", "utf8", str(Path(__file__).resolve()), "case", case.name,
                   "--out", str(target), "--state", str(state / case.name)]
        env["THERMOGAR_STATE_ROOT"] = str(state / case.name / "state")
        log_path = out.parent / f"{out.name}_logs" / f"{case.name}.log"
        log_path.parent.mkdir(parents=True, exist_ok=True)
        stamp = dt.datetime.now().astimezone().isoformat(timespec="seconds")
        started = time.perf_counter()
        peak = 0
        min_seen = free
        aborted = False
        with log_path.open("wb") as log:
            process = subprocess.Popen(command, cwd=PROJECT_ROOT, env=env, stdout=log, stderr=subprocess.STDOUT)
            parent = psutil.Process(process.pid)
            while process.poll() is None:
                total = 0
                try:
                    for member in [parent] + parent.children(recursive=True):
                        try:
                            total += member.memory_info().rss
                        except psutil.Error:
                            pass
                except psutil.Error:
                    pass
                peak = max(peak, total)
                now_free = psutil.virtual_memory().available / GIB
                min_seen = min(min_seen, now_free)
                if now_free < ABORT_FREE_GIB:
                    aborted = True
                    for member in parent.children(recursive=True) + [parent]:
                        try:
                            member.kill()
                        except psutil.Error:
                            pass
                    break
                time.sleep(0.5)
            code = process.wait()
        seconds = time.perf_counter() - started
        writer.writerow([case.name, case.group, case.slow, stamp, f"{seconds:.1f}", f"{peak / GIB:.2f}",
                         f"{free:.2f}", f"{min_seen:.2f}", f"{waited:.0f}", code, aborted])
        handle.flush()
        elapsed = time.perf_counter() - run_started
        status = "снят" if code == 0 else f"КОД {code}"
        print(f"[{number}/{len(cases)}] {case.name}: {status}, {seconds:.1f} с, пик {peak / GIB:.2f} ГиБ, "
              f"всего {elapsed / 60:.1f} мин", flush=True)
        if code != 0:
            failures += 1
        if aborted:
            print("СТОП: снят по памяти", flush=True)
            return 4
    handle.close()
    return 1 if failures else 0


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    sub = parser.add_subparsers(dest="command", required=True)
    sub.add_parser("list")
    one = sub.add_parser("case")
    one.add_argument("name")
    one.add_argument("--out", required=True, type=Path)
    one.add_argument("--state", required=True, type=Path)
    every = sub.add_parser("run")
    every.add_argument("--out", required=True, type=Path)
    every.add_argument("--state", required=True, type=Path)
    every.add_argument("--time-csv", required=True, type=Path)
    every.add_argument("--only", default=None, help="регулярное выражение по имени случая")
    every.add_argument("--min-free", type=float, default=3.0, help="порог входа, ГиБ свободной памяти")
    every.add_argument("--wait", type=float, default=1800.0, help="сколько ждать памяти на входе, с")
    args = parser.parse_args()
    if args.command == "list":
        for case in build_cases():
            print(f"{case.group} | {case.name} | {'slow' if case.slow else '-'} | {case.source}")
        return 0
    if args.command == "case":
        return run_case(args.name, args.out.resolve(), args.state.resolve())
    return run_all(args.out.resolve(), args.state.resolve(), args.time_csv.resolve(), args.only,
                   args.min_free, args.wait)


if __name__ == "__main__":
    raise SystemExit(main())
